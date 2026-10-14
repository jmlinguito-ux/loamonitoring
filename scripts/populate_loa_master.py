"""Populate Dataverse LOA Monitoring table (cr63f_loamonitoring)

Imports records from 'LOA (NEW) (1).xlsx' (Sheet: Append1) into Dataverse using
high-performance OData v9.2 multipart $batch requests with automatic checkpointing,
deduplication, date parsing, and exponential throttling backoff.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
import uuid
from datetime import datetime, timedelta
from typing import Any, Dict, List, Optional, Set

import openpyxl
import requests

# Ensure parent directory is in path so we can import dataverse module
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, '..'))
if PROJECT_DIR not in sys.path:
    sys.path.insert(0, PROJECT_DIR)

import dataverse as dv


CHECKPOINT_DIR = SCRIPT_DIR
IMPORTED_LIST_FILE = os.path.join(CHECKPOINT_DIR, "imported_controls.txt")
CHECKPOINT_STATE_FILE = os.path.join(CHECKPOINT_DIR, "import_checkpoint.json")


def parse_num(val: Any) -> Optional[float]:
    """Parse numeric values (currencies, areas, percentages)."""
    if val is None or isinstance(val, bool):
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s or s.lower() in ('none', 'null', 'nan', '-', '#n/a', 'n/a', 'undefined'):
        return None
    s = re.sub(r'[₱$\s,]', '', s)
    if s.endswith('%'):
        s = s[:-1]
    if s.startswith('(') and s.endswith(')'):
        s = '-' + s[1:-1]
    try:
        return float(s)
    except ValueError:
        return None


def parse_int(val: Any) -> Optional[int]:
    """Parse integer values (such as INDEX NO)."""
    num = parse_num(val)
    if num is None:
        return None
    try:
        return int(round(num))
    except (ValueError, OverflowError):
        return None


def parse_date(val: Any) -> Optional[str]:
    """Parse date values into ISO 8601 UTC strings ('YYYY-MM-DDTHH:MM:SSZ')."""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime('%Y-%m-%dT00:00:00Z')
    if isinstance(val, (int, float)):
        if val > 1000:
            try:
                dt = datetime(1899, 12, 30) + timedelta(days=val)
                if 1753 <= dt.year <= 9999:
                    return dt.strftime('%Y-%m-%dT00:00:00Z')
            except Exception:
                pass
        return None
    s = str(val).strip()
    if not s or s.lower() in ('none', 'null', 'nan', '-', '0', '1', 'undefined', 'n/a'):
        return None
    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%d/%m/%Y', '%Y/%m/%d', '%m/%d/%y'):
        try:
            dt = datetime.strptime(s, fmt)
            if 1753 <= dt.year <= 9999:
                return dt.strftime('%Y-%m-%dT00:00:00Z')
        except ValueError:
            continue
    return None


def clean_str(val: Any, max_len: Optional[int] = None) -> Optional[str]:
    """Clean string values, returning None if empty or placeholder."""
    if val is None:
        return None
    s = str(val).strip()
    if not s or s.lower() in ('none', 'null', 'nan', '-', 'undefined', 'n/a', '#n/a'):
        return None
    if max_len and len(s) > max_len:
        s = s[:max_len]
    return s


def load_imported_controls() -> Set[str]:
    """Load previously imported control numbers from checkpoint."""
    if not os.path.exists(IMPORTED_LIST_FILE):
        return set()
    with open(IMPORTED_LIST_FILE, 'r', encoding='utf-8') as f:
        return {line.strip() for line in f if line.strip()}


def append_imported_controls(controls: List[str]) -> None:
    """Append batch of imported controls to checkpoint file."""
    with open(IMPORTED_LIST_FILE, 'a', encoding='utf-8') as f:
        for c in controls:
            f.write(c + '\n')


def update_checkpoint_state(total_records: int, imported_count: int, failed_controls: List[str]) -> None:
    """Save high-level import state to JSON."""
    data = {
        'total_records': total_records,
        'imported_count': imported_count,
        'updated_at': datetime.now().isoformat(),
        'failed_controls': failed_controls,
    }
    with open(CHECKPOINT_STATE_FILE, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=2)


def read_and_deduplicate_excel(excel_path: str, sheet_name: str) -> List[Dict[str, Any]]:
    """Read Excel sheet and normalize records, resolving duplicates on CONTROL NO."""
    print(f"[{datetime.now().strftime('%H:%M:%S')}] Loading '{excel_path}' (sheet: {sheet_name})...", flush=True)
    t0 = time.time()
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found. Available sheets: {wb.sheetnames}")
    ws = wb[sheet_name]
    row_iter = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(next(row_iter))]

    records_by_ctrl: Dict[str, Dict[str, Any]] = {}
    total_raw_rows = 0
    skipped_no_ctrl = 0
    duplicates_resolved = 0

    for row in row_iter:
        total_raw_rows += 1
        raw = dict(zip(headers, row))
        ctrl = raw.get('CONTROL NO')
        if ctrl is None:
            skipped_no_ctrl += 1
            continue
        ctrl_str = str(ctrl).strip()
        if not ctrl_str or ctrl_str.lower() in ('none', 'null', 'nan', 'control no', '-', ''):
            skipped_no_ctrl += 1
            continue

        if ctrl_str in records_by_ctrl:
            duplicates_resolved += 1
            existing = records_by_ctrl[ctrl_str]
            # Prioritize rows with populated TYPE or merge non-empty fields
            if not existing.get('TYPE') and raw.get('TYPE'):
                for k, v in raw.items():
                    if v is not None and str(v).strip() != '':
                        existing[k] = v
            else:
                for k, v in raw.items():
                    if (existing.get(k) is None or str(existing.get(k)).strip() == '') and v is not None and str(v).strip() != '':
                        existing[k] = v
        else:
            records_by_ctrl[ctrl_str] = raw

    wb.close()
    read_time = time.time() - t0
    print(f"[{datetime.now().strftime('%H:%M:%S')}] Read {total_raw_rows} rows in {read_time:.2f}s.", flush=True)
    print(f"  - Skipped blank/invalid CONTROL NO rows: {skipped_no_ctrl}", flush=True)
    print(f"  - Duplicate CONTROL NO rows merged: {duplicates_resolved}", flush=True)
    print(f"  - Unique records to process: {len(records_by_ctrl):,}", flush=True)

    # Convert to Dataverse payload dictionaries
    payloads: List[Dict[str, Any]] = []
    for ctrl, raw in records_by_ctrl.items():
        lot_area = parse_num(raw.get('LOT AREA (HA)'))
        rcc_price = parse_num(raw.get('RCC PRICE'))
        tcp = parse_num(raw.get('TOTAL CONTRACT PRICE'))
        deductible = parse_num(raw.get('TOTAL PAYMENT DEDUCTIBLE TO TCP'))
        balance = parse_num(raw.get('BALANCE'))
        paid_raw = parse_num(raw.get('% PAID'))

        # Calculate derived financial figures if appropriate
        if tcp is not None and tcp > 0:
            if deductible is not None:
                if balance is None:
                    balance = max(0.0, tcp - deductible)
                if paid_raw is None:
                    paid_raw = (deductible / tcp) * 100.0

        paid_pct = None
        if paid_raw is not None:
            if 0.0 < paid_raw <= 1.0:
                paid_pct = round(paid_raw * 100.0, 1)
            else:
                paid_pct = round(paid_raw, 1)

        p = {
            'cr63f_control1': ctrl,
            'cr63f_type': clean_str(raw.get('TYPE')),
            'cr63f_province': clean_str(raw.get('PROVINCE')),
            'cr63f_municipality': clean_str(raw.get('MUNICIPALITY')),
            'cr63f_dbadjustedbarangay': clean_str(raw.get('DB ADJUSTED BARANGAY')),
            'cr63f_indexno': parse_int(raw.get('INDEX NO')),
            'cr63f_areaindex': clean_str(raw.get('AREA-INDEX')),
            'cr63f_registeredowner': clean_str(raw.get('REGISTERED OWNER')),
            'cr63f_authorizedrepresentative': clean_str(raw.get('AUTHORIZED REPRESENTATIVE')),
            'cr63f_lotareaha': round(lot_area, 4) if lot_area is not None else None,
            'cr63f_titleno': clean_str(raw.get('TITLE NO')),
            'new_mode': clean_str(raw.get('MODE')),
            'cr63f_acquisitionsublevel': clean_str(raw.get('ACQUISITION LEVEL')),
            'cr63f_rccprice': round(rcc_price, 2) if rcc_price is not None else None,
            'cr63f_rccterms': clean_str(raw.get('RCC TERMS')),
            'cr63f_entity': clean_str(raw.get('ENTITY')),
            'new_contracttype': clean_str(raw.get('CONTRACT TYPE')),
            'cr63f_conversionstatus': clean_str(raw.get('CONVERSION STATUS')),
            'cr63f_totalcontractprice': round(tcp, 2) if tcp is not None else None,
            'cr63f_totalpaymentdeductibletotcp': round(deductible, 2) if deductible is not None else None,
            'cr63f_paid': paid_pct,
            'cr63f_balance': round(balance, 2) if balance is not None else None,
            'cr63f_securedendorsedate': parse_date(raw.get('SECURED ENDORSE DATE')),
            'cr63f_endorseddate': parse_date(raw.get('ENDORSED DATE')),
        }
        clean_p = {k: v for k, v in p.items() if v is not None}
        payloads.append(clean_p)

    return payloads


def ensure_token(client: dv.DataverseClient, max_retries: int = 5) -> str:
    """Ensure client has a valid token, retrying silently if a transient network blip occurred."""
    for attempt in range(max_retries):
        try:
            return client.token()
        except Exception:
            if attempt == max_retries - 1:
                raise
            time.sleep(2.0 * (attempt + 1))
            client._token = None
            client._token_expiry = 0
            if client.try_silent_token():
                return client._token  # type: ignore
    raise PermissionError("Failed to refresh Dataverse token.")


def send_batch_with_retry(
    client: dv.DataverseClient,
    entity_set: str,
    batch_payloads: List[Dict[str, Any]],
    max_retries: int = 5
) -> Tuple[List[str], List[Tuple[str, str]]]:
    """Send a multipart $batch request with automatic retry on 429 / 503.
    
    Returns (list_of_succeeded_controls, list_of_failed_controls_with_error).
    """
    batch_id = 'batch_' + uuid.uuid4().hex
    lines = []
    for n, p in enumerate(batch_payloads, start=1):
        cs_id = 'changeset_' + uuid.uuid4().hex
        lines.extend([
            '--' + batch_id,
            'Content-Type: multipart/mixed; boundary=' + cs_id,
            '',
            '--' + cs_id,
            'Content-Type: application/http',
            'Content-Transfer-Encoding: binary',
            'Content-ID: ' + str(n),
            '',
            'POST ' + client.api_root + entity_set + ' HTTP/1.1',
            'Content-Type: application/json;type=entry',
            '',
            json.dumps(p),
            '--' + cs_id + '--'
        ])
    lines.append('--' + batch_id + '--')
    body = '\r\n'.join(lines).encode('utf-8')

    url = client.api_root + '$batch'

    for attempt in range(max_retries):
        try:
            token = ensure_token(client)
            headers = {
                'Authorization': f'Bearer {token}',
                'Accept': 'application/json',
                'OData-MaxVersion': '4.0',
                'OData-Version': '4.0',
                'Content-Type': f'multipart/mixed; boundary={batch_id}'
            }

            resp = client._session.post(url, headers=headers, data=body, timeout=120)

            if resp.status_code in (429, 503):
                try:
                    retry_after = float(resp.headers.get('Retry-After', 0))
                except (ValueError, TypeError):
                    retry_after = 0
                wait_sec = max(retry_after, 2.0 * (2 ** attempt))
                print(f"  [Throttle {resp.status_code}] Waiting {wait_sec:.1f}s before retry {attempt + 1}/{max_retries}...", flush=True)
                time.sleep(wait_sec)
                continue

            if resp.status_code == 200:
                succeeded: List[str] = []
                failed: List[Tuple[str, str]] = []

                # Split by boundary lines to parse each changeset response
                parts = re.split(r'--changesetresponse_[0-9a-fA-F-]+', resp.text)
                for part in parts:
                    cid_match = re.search(r'Content-ID:\s*(\d+)', part)
                    status_match = re.search(r'HTTP/1\.1\s+(\d{3})', part)
                    if cid_match and status_match:
                        idx = int(cid_match.group(1)) - 1
                        if 0 <= idx < len(batch_payloads):
                            ctrl = batch_payloads[idx]['cr63f_control1']
                            status_code = int(status_match.group(1))
                            if 200 <= status_code < 300:
                                succeeded.append(ctrl)
                            else:
                                err_match = re.search(r'"message":\s*"([^"]+)"', part)
                                err_msg = err_match.group(1) if err_match else f"HTTP {status_code}"
                                failed.append((ctrl, err_msg))

                # If parts parsing didn't find any Content-IDs (unexpected format fallback)
                if not succeeded and not failed:
                    created_count = len(re.findall(r'HTTP/1\.1 (?:204|201|200)', resp.text))
                    if created_count == len(batch_payloads):
                        return [p['cr63f_control1'] for p in batch_payloads], []
                    else:
                        return [], [(p['cr63f_control1'], "Parse mismatch") for p in batch_payloads]

                return succeeded, failed
            else:
                err_text = f"HTTP {resp.status_code}: {resp.text[:200]}"
                if attempt == max_retries - 1:
                    return [], [(p['cr63f_control1'], err_text) for p in batch_payloads]
                time.sleep(2.0 * (2 ** attempt))

        except (requests.exceptions.ConnectionError, requests.exceptions.Timeout) as ex:
            if attempt == max_retries - 1:
                return [], [(p['cr63f_control1'], str(ex)) for p in batch_payloads]
            wait_sec = 2.0 * (2 ** attempt)
            print(f"  [Network Error] {ex}. Waiting {wait_sec:.1f}s before retry...", flush=True)
            time.sleep(wait_sec)

    return [], [(p['cr63f_control1'], "Max retries exceeded") for p in batch_payloads]


def main():
    parser = argparse.ArgumentParser(description="Populate Dataverse LOA Monitoring table from Excel")
    parser.add_argument("--excel", default=r"C:\Users\Jinguito\Downloads\LOA (NEW) (1).xlsx", help="Excel file path")
    parser.add_argument("--sheet", default="Append1", help="Sheet name in Excel")
    parser.add_argument("--batch-size", type=int, default=100, help="Batch size for Dataverse OData $batch")
    parser.add_argument("--limit", type=int, default=None, help="Limit number of records to process")
    parser.add_argument("--reset", action="store_true", help="Reset checkpoint and re-import from scratch")
    args = parser.parse_args()

    print("=" * 70, flush=True)
    print("LOA Monitoring Dataverse Ingestion Pipeline", flush=True)
    print("=" * 70, flush=True)

    if args.reset:
        if os.path.exists(IMPORTED_LIST_FILE):
            os.remove(IMPORTED_LIST_FILE)
        if os.path.exists(CHECKPOINT_STATE_FILE):
            os.remove(CHECKPOINT_STATE_FILE)
        print("Checkpoint reset successfully.", flush=True)

    # 1. Connect to Dataverse
    client = dv.DataverseClient()
    if not client.signed_in():
        print("ERROR: Not signed in to Dataverse! Please run app.py or login first.", flush=True)
        sys.exit(1)
    user_email = client.current_user_name()
    print(f"Authenticated as: {user_email}", flush=True)
    print(f"Environment: {client.env}", flush=True)

    info = client.get_entity_info('cr63f_loamonitoring')
    entity_set = info['entity_set']
    print(f"Target Entity Set: {entity_set}", flush=True)

    # 2. Read and deduplicate Excel
    all_records = read_and_deduplicate_excel(args.excel, args.sheet)
    if args.limit:
        all_records = all_records[:args.limit]
        print(f"Applied limit: processing {len(all_records)} records.", flush=True)

    total_records = len(all_records)

    # 3. Filter already imported records from checkpoint
    already_imported = load_imported_controls()
    pending_records = [r for r in all_records if r['cr63f_control1'] not in already_imported]

    print(f"\nTotal unique records: {total_records:,}", flush=True)
    print(f"Already imported in previous runs: {len(already_imported):,}", flush=True)
    print(f"Records remaining to import: {len(pending_records):,}\n", flush=True)

    if not pending_records:
        print("All records are already imported! Nothing to do.", flush=True)
        return

    # 4. Batch ingestion loop
    batch_size = args.batch_size
    total_batches = (len(pending_records) + batch_size - 1) // batch_size
    failed_controls: List[str] = []
    imported_in_this_run = 0

    start_time = time.time()
    print(f"Starting ingestion: {len(pending_records):,} records in {total_batches:,} batches of {batch_size}...", flush=True)

    for b_idx in range(total_batches):
        chunk = pending_records[b_idx * batch_size : (b_idx + 1) * batch_size]
        t_batch_start = time.time()

        succeeded_ctrls, failed_batch = send_batch_with_retry(client, entity_set, chunk)

        if succeeded_ctrls:
            append_imported_controls(succeeded_ctrls)
            imported_in_this_run += len(succeeded_ctrls)

        if failed_batch:
            for failed_c, failed_err in failed_batch:
                print(f"  [ERROR] Record {failed_c} failed: {failed_err}", flush=True)
                failed_controls.append(failed_c)

        total_done = len(already_imported) + imported_in_this_run
        pct_done = (total_done / total_records) * 100.0

        elapsed = time.time() - start_time
        rate = imported_in_this_run / elapsed if elapsed > 0 else 0
        remaining_records = total_records - total_done
        eta_sec = (remaining_records / rate) if rate > 0 else 0
        eta_min = eta_sec / 60.0

        batch_time = time.time() - t_batch_start
        now_str = datetime.now().strftime('%H:%M:%S')
        print(
            f"[{now_str}] Batch {b_idx + 1:>3}/{total_batches} ({len(chunk):>3} rows in {batch_time:>4.2f}s) | "
            f"Total: {total_done:>5,}/{total_records:,} ({pct_done:>5.1f}%) | "
            f"Speed: {rate:>4.1f} rows/s | ETA: {eta_min:>4.1f}m",
            flush=True
        )

        if (b_idx + 1) % 5 == 0 or (b_idx + 1) == total_batches:
            update_checkpoint_state(total_records, total_done, failed_controls)

    total_time = time.time() - start_time
    print("\n" + "=" * 70, flush=True)
    print("Ingestion Finished!", flush=True)
    print(f"Total time elapsed: {total_time / 60.0:.2f} minutes", flush=True)
    print(f"Records imported in this run: {imported_in_this_run:,}", flush=True)
    print(f"Total imported records: {len(already_imported) + imported_in_this_run:,} / {total_records:,}", flush=True)
    if failed_controls:
        print(f"WARNING: {len(failed_controls)} records failed to import. Check {CHECKPOINT_STATE_FILE}", flush=True)
    print("=" * 70, flush=True)


if __name__ == "__main__":
    main()
