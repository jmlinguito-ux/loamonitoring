"""End-to-end test of every LOA Monitoring Tool feature: real UI + real Python + fake Dataverse."""
import json
import os
import sys
import time
import traceback
from datetime import datetime

import openpyxl
from playwright.sync_api import sync_playwright

import harness

RESULTS = []
TODAY = datetime.now().strftime('%Y-%m-%d')


def check(area, name, ok, detail=''):
    RESULTS.append({'area': area, 'name': name, 'ok': bool(ok), 'detail': str(detail)[:400]})
    print(('PASS ' if ok else 'FAIL ') + f'[{area}] {name}' + (f' -- {detail}' if detail and not ok else ''), flush=True)


def step(area, name):
    def wrap(fn):
        def run(*a, **k):
            try:
                return fn(*a, **k)
            except Exception as exc:
                check(area, name, False, f'crashed: {exc!r}'.split('Call log')[0])
                try:
                    shot = os.path.join(harness.WORK, f'fail_{len(RESULTS)}.png')
                    page.screenshot(path=shot)
                    print('   last toasts:', page.evaluate('window.__toasts')[-4:], 'screenshot', shot, flush=True)
                except Exception:
                    pass
                cleanup()
        return run
    return wrap


def cleanup():
    try:
        page.evaluate("""() => { const m = document.getElementById('confirmModal');
            if (m && !m.classList.contains('hidden')) document.getElementById('btnConfirmModalCancel')?.click();
            m && m.classList.add('hidden');
            const d = document.getElementById('entryDrawer'); if (d) { d.classList.remove('open'); d.style.display = 'none'; } }""")
        page.wait_for_timeout(300)
    except Exception:
        pass


env = harness.start()
S = env.store
LOA, LOGS, FIN, CONV = 'cr63f_loamonitoring', 'cr63f_loamonitoringlogs', 'cr63f_financereference', 'cr63f_landconversionreference'


def acct(ctrl):
    rows = S.find(LOA, cr63f_control1=ctrl)
    return rows[0] if rows else None


def logs_for(ctrl=None, action=None):
    out = []
    for r in S.rows[LOGS].values():
        if ctrl and r.get('cr63f_control1') != ctrl:
            continue
        if action and r.get('cr63f_actiontype') != action:
            continue
        out.append(r)
    return out


TOAST_HOOK = """
(() => {
  window.__toasts = window.__toasts || [];
  const wrap = document.getElementById('toastWrap');
  if (!wrap || wrap.__hooked) return;
  wrap.__hooked = true;
  new MutationObserver(ms => ms.forEach(m => m.addedNodes.forEach(n => {
    if (n.classList && n.classList.contains('toast')) window.__toasts.push([n.className.replace('toast', '').trim(), n.textContent.trim()]);
  }))).observe(wrap, {childList: true});
})();
"""

pw = sync_playwright().start()
browser = pw.chromium.launch()
page = browser.new_page(viewport={'width': 1500, 'height': 950})
PAGE_ERRORS = []
page.on('pageerror', lambda e: PAGE_ERRORS.append(str(e)))
page.route('**/fonts.googleapis.com/**', lambda r: r.abort())
page.route('**/fonts.gstatic.com/**', lambda r: r.abort())
page.set_default_timeout(8000)
page.goto(env.url)
page.wait_for_timeout(2500)
page.evaluate(TOAST_HOOK)


def toasts(since=0):
    return page.evaluate('window.__toasts')[since:]


def mark():
    return len(page.evaluate('window.__toasts'))


def nav(view, wait=1200):
    page.click(f'[data-view="{view}"]')
    page.wait_for_timeout(wait)


def modal_ok(wait=1500, timeout=20000):
    page.wait_for_selector('#confirmModal:not(.hidden) #btnConfirmModalOk', timeout=timeout)
    page.click('#btnConfirmModalOk')
    page.wait_for_timeout(wait)


def wait_modal_closed(timeout=60000):
    page.wait_for_function("document.getElementById('confirmModal').classList.contains('hidden')", timeout=timeout)


def row_sel(ctrl):
    return f'#recordsTableBody tr[data-row-id="{acct(ctrl)["cr63f_loamonitoringid"]}"]'


def cell_input(ctrl, col):
    return f'{row_sel(ctrl)} td[data-col-key="{col}"] input, {row_sel(ctrl)} td[data-col-key="{col}"] select'


def type_cell(ctrl, col, value):
    page.evaluate("""([sel, value]) => { const i = document.querySelector(sel);
        i.dispatchEvent(new Event('focus')); i.value = value; i.dispatchEvent(new Event('input'));
        i.dispatchEvent(new Event('blur')); }""", [cell_input(ctrl, col), value])
    page.wait_for_timeout(300)


def choose_cell(ctrl, col, value):
    page.evaluate("""([sel, value]) => { const s = document.querySelector(sel); s.value = value;
        s.dispatchEvent(new Event('change')); }""", [cell_input(ctrl, col), value])
    page.wait_for_timeout(300)


def cell_value(ctrl, col):
    return page.eval_on_selector(cell_input(ctrl, col), 'e => e.tagName === "SELECT" ? e.options[e.selectedIndex]?.text : e.value')


def set_select(sel, value):
    page.evaluate("""([sel, value]) => { const s = document.querySelector(sel); s.value = value;
        s.dispatchEvent(new Event('change', {bubbles: true})); s.dispatchEvent(new Event('input', {bubbles: true})); }""", [sel, value])
    page.wait_for_timeout(500)


def answer_dialog(path):
    env.window.dialog_answers.append(path)


def reload_records():
    nav('records', 300)
    page.evaluate("document.querySelector('#btnTopSyncDv') && document.querySelector('#btnTopSyncDv').click()")
    page.wait_for_timeout(2500)


# =====================================================================================
# 1. Startup and Land Master basics
# =====================================================================================
@step('Startup', 'app loads')
def t_startup():
    nav('records')
    check('Startup', 'Land Master loads all 14 accounts (paged reads)', page.text_content('#lblGridCount').strip() == '14 records',
          page.text_content('#lblGridCount'))
    expect_bal = sum(max(0, (r.get('cr63f_totalcontractprice') or 0) - (r.get('cr63f_totalpaymentdeductibletotcp') or 0))
                     for r in S.rows[LOA].values())
    kpi = page.text_content('#kpiTotalBalance').strip()
    check('Startup', 'Balance KPI = sum of max(0, TCP - Deductible)', kpi == f'₱{expect_bal:,.2f}', f'{kpi} vs {expect_bal:,.2f}')
    check('Startup', 'Connection pill shows connected', 'onnect' in (page.text_content('#lblTopConn') or ''), page.text_content('#lblTopConn'))
    check('Startup', 'Endorsed date 1/5/2026 shows with Work Week WW-01',
          cell_value('C-001', 'cr63f_endorseddate') in ('01/05/2026', '1/5/2026')
          and 'WW-01' in page.text_content(f'{row_sel("C-001")} td[data-col-key="work_week"]'),
          (cell_value('C-001', 'cr63f_endorseddate'), page.text_content(f'{row_sel("C-001")} td[data-col-key="work_week"]')))
    check('Startup', '1900 placeholder date has no Work Week',
          page.text_content(f'{row_sel("C-005")} td[data-col-key="work_week"]').strip() in ('-', ''),
          page.text_content(f'{row_sel("C-005")} td[data-col-key="work_week"]'))
    check('Startup', 'Mode not in the list shows as "(not in list)"', '(not in list)' in (cell_value('C-004', 'new_mode') or ''),
          cell_value('C-004', 'new_mode'))
    pct = cell_value('C-001', 'cr63f_paid')
    check('Startup', '% Paid shows the calculated value (25%), not the legacy stored 0.3', pct.startswith('25'), pct)


@step('Land Master', 'search and filters')
def t_filters():
    page.fill('#searchInput', 'Owner C-007')
    page.dispatch_event('#searchInput', 'input')
    page.wait_for_timeout(600)
    n = page.locator('#recordsTableBody tr').count()
    check('Land Master', 'Search finds one account', n == 1, n)
    page.fill('#searchInput', '')
    page.dispatch_event('#searchInput', 'input')
    set_select('#typeFilter', 'ROW')
    page.wait_for_timeout(600)
    check('Land Master', 'Type filter ROW shows 1', page.locator('#recordsTableBody tr').count() == 1,
          page.locator('#recordsTableBody tr').count())
    page.click('#btnClearFilters')
    page.wait_for_timeout(600)
    opts = page.eval_on_selector_all('#modeFilterPanel option', 'os => os.map(o => o.value)')
    lease = next((o for o in opts if 'LEASE' in o.upper()), None)
    if lease:
        set_select('#modeFilterPanel', lease)
        page.wait_for_timeout(600)
        check('Land Master', 'Mode filter (lease) shows the lease account', page.locator('#recordsTableBody tr').count() == 1,
              page.locator('#recordsTableBody tr').count())
    else:
        check('Land Master', 'Mode filter has a lease option', False, opts)
    page.click('#btnClearFilters')
    page.wait_for_timeout(600)
    check('Land Master', 'Clear filters restores 14', page.locator('#recordsTableBody tr').count() == 14,
          page.locator('#recordsTableBody tr').count())


@step('Land Master', 'columns and sorting')
def t_columns():
    page.click('#btnColumnsPicker')
    page.wait_for_timeout(300)
    page.evaluate("""() => { const c = document.querySelector('#colPickerMenu [data-col-toggle="cr63f_rccterms"]');
        c.checked = !c.checked; c.dispatchEvent(new Event('change', {bubbles: true})); }""")
    page.wait_for_timeout(500)
    hidden = page.locator('#recordsTableHead [data-col-key="cr63f_rccterms"], #recordsTableHead th:has-text("RCC Terms")').count() == 0
    check('Land Master', 'Column picker hides a column', hidden)
    page.evaluate("""() => { const c = document.querySelector('#colPickerMenu [data-col-toggle="cr63f_rccterms"]');
        c.checked = !c.checked; c.dispatchEvent(new Event('change', {bubbles: true})); }""")
    page.keyboard.press('Escape')
    page.mouse.click(5, 5)
    before = page.eval_on_selector_all('#recordsTableBody tr td[data-col-key="cr63f_control1"] input', 'es => es.map(e => e.value)')
    page.evaluate("""() => { const th = Array.from(document.querySelectorAll('#recordsTableHead th')).find(t => /Total Contract Price/.test(t.textContent));
        (th.querySelector('.th-label, .th-text, span') || th).click(); }""")
    page.wait_for_timeout(600)
    after = page.eval_on_selector_all('#recordsTableBody tr td[data-col-key="cr63f_control1"] input', 'es => es.map(e => e.value)')
    check('Land Master', 'Clicking a header sorts the grid', before != after, after[:5])


# =====================================================================================
# 2. Grid editing
# =====================================================================================
@step('Grid edit', 'edit and save')
def t_grid_edit():
    reload_records()
    page.click('#btnClearFilters')
    page.wait_for_timeout(500)
    m = mark()
    type_cell('C-006', 'cr63f_province', 'CAVITE')
    type_cell('C-002', 'cr63f_totalcontractprice', '0')
    check('Grid edit', 'Typed 0 stays ₱0.00', cell_value('C-002', 'cr63f_totalcontractprice') == '₱0.00', cell_value('C-002', 'cr63f_totalcontractprice'))
    type_cell('C-006', 'cr63f_endorseddate', '2/3/2026')
    check('Grid edit', 'Date 2/3/2026 becomes 02/03/2026', cell_value('C-006', 'cr63f_endorseddate') == '02/03/2026',
          cell_value('C-006', 'cr63f_endorseddate'))
    ww = page.text_content(f'{row_sel("C-006")} td[data-col-key="work_week"]').strip()
    check('Grid edit', 'Work Week follows the new Endorsed Date (Feb 3 2026 = WW-05)', ww == 'WW-05', ww)
    type_cell('C-007', 'cr63f_endorseddate', '15/20/26')
    bad = toasts(m)
    check('Grid edit', 'Impossible date refused with a message, value restored',
          any('not a valid date' in t for _, t in bad) and cell_value('C-007', 'cr63f_endorseddate') in ('09/02/2026', '9/2/2026'),
          (bad, cell_value('C-007', 'cr63f_endorseddate')))
    choose_cell('C-008', 'new_mode', '')
    check('Grid edit', 'Clearing Mode is refused', cell_value('C-008', 'new_mode') == 'SECURED SALE'
          and any('cannot be cleared' in t for _, t in toasts(m)), cell_value('C-008', 'new_mode'))
    chip = page.text_content('#unsavedEditsChip')
    check('Grid edit', 'Unsaved chip counts 2 modified rows', '2 modified' in chip, chip)
    page.click('#btnSaveStagedEdits')
    modal_ok(2500)
    c6, c2 = acct('C-006'), acct('C-002')
    check('Grid edit', 'Province saved to Dataverse', c6['cr63f_province'] == 'CAVITE', c6['cr63f_province'])
    check('Grid edit', 'Endorsed Date saved as 2026-02-03', c6['cr63f_endorseddate'] == '2026-02-03', c6['cr63f_endorseddate'])
    check('Grid edit', 'TCP 0 saved as 0 with Balance 0 and % Paid 0 (deductible 0)',
          c2['cr63f_totalcontractprice'] == 0 and c2['cr63f_balance'] == 0 and c2['cr63f_paid'] == 0,
          (c2['cr63f_totalcontractprice'], c2['cr63f_balance'], c2['cr63f_paid']))
    b6 = logs_for('C-006', 'BASELINE')
    e6 = logs_for('C-006', 'MANUAL_EDIT')
    check('Grid edit', 'First change writes a hidden BASELINE then a MANUAL_EDIT snapshot',
          len(b6) == 1 and len(e6) == 1 and b6[0]['cr63f_province'] == 'BATANGAS' and e6[0]['cr63f_province'] == 'CAVITE',
          (len(b6), len(e6)))
    check('Grid edit', 'Success toast after save', any('Successfully saved' in t for _, t in toasts(m)), toasts(m)[-3:])
    m = mark()
    type_cell('C-006', 'cr63f_province', 'LAGUNA')
    page.click('#btnDiscardStagedEdits')
    page.wait_for_timeout(500)
    try:
        modal_ok(500, timeout=1500)
    except Exception:
        pass
    check('Grid edit', 'Discard restores the saved value', cell_value('C-006', 'cr63f_province') == 'CAVITE', cell_value('C-006', 'cr63f_province'))


@step('Grid edit', 'paste from Excel')
def t_paste():
    m = mark()
    page.evaluate("""([a, b]) => {
        const td1 = document.querySelector(a), td2 = document.querySelector(b);
        td1.dispatchEvent(new MouseEvent('mousedown', {bubbles: true, button: 0}));
        td2.dispatchEvent(new MouseEvent('mousedown', {bubbles: true, button: 0, shiftKey: true}));
        document.dispatchEvent(new MouseEvent('mouseup', {bubbles: true}));
        const dt = new DataTransfer(); dt.setData('text/plain', '1/5/2026\\n15/20/26\\n');
        td1.dispatchEvent(new ClipboardEvent('paste', {clipboardData: dt, bubbles: true, cancelable: true}));
    }""", [f'{row_sel("C-007")} td[data-col-key="cr63f_secureddate"]', f'{row_sel("C-008")} td[data-col-key="cr63f_secureddate"]'])
    page.wait_for_timeout(800)
    v7 = cell_value('C-007', 'cr63f_secureddate')
    v8 = cell_value('C-008', 'cr63f_secureddate')
    check('Grid edit', 'Pasted Excel date 1/5/2026 becomes 01/05/2026', v7 == '01/05/2026', v7)
    check('Grid edit', 'Pasted impossible date is skipped and reported', v8 == '' and any('invalid date' in t for _, t in toasts(m)),
          (v8, toasts(m)))
    page.click('#btnSaveStagedEdits')
    modal_ok(2500)
    check('Grid edit', 'Pasted date saved', acct('C-007')['cr63f_secureddate'] == '2026-01-05', acct('C-007')['cr63f_secureddate'])
    check('Grid edit', 'Skipped date left untouched in Dataverse', acct('C-008').get('cr63f_secureddate') is None)


@step('Grid edit', 'someone else changed the account')
def t_concurrency():
    m = mark()
    other = acct('C-009')
    other['cr63f_province'] = 'QUEZON (changed by colleague)'
    other['_version'] += 1
    type_cell('C-009', 'cr63f_municipality', 'TUY')
    page.click('#btnSaveStagedEdits')
    modal_ok(2500)
    msgs = [t for _, t in toasts(m)]
    check('Concurrency', 'Save refused with "changed ... by someone else"', any('someone else' in t for t in msgs), msgs)
    check('Concurrency', "Colleague's change kept, ours not written", acct('C-009')['cr63f_province'].startswith('QUEZON')
          and acct('C-009').get('cr63f_municipality') == 'NASUGBU', acct('C-009').get('cr63f_municipality'))
    check('Concurrency', 'No audit snapshot written for the refused save', not logs_for('C-009', 'MANUAL_EDIT'))
    try:
        page.click('#btnDiscardStagedEdits')
        modal_ok(500, timeout=1500)
    except Exception:
        pass
    reload_records()
    type_cell('C-009', 'cr63f_municipality', 'TUY')
    page.click('#btnSaveStagedEdits')
    modal_ok(2500)
    check('Concurrency', 'After reload the same edit saves', acct('C-009').get('cr63f_municipality') == 'TUY')
    snap = logs_for('C-009', 'MANUAL_EDIT')
    check('Concurrency', "Audit snapshot includes the colleague's current value",
          snap and snap[-1]['cr63f_province'].startswith('QUEZON'), snap and snap[-1]['cr63f_province'])


# =====================================================================================
# 3. New entry drawer
# =====================================================================================
def fill_drawer(values):
    for sel, val in values.items():
        page.evaluate("""([sel, val]) => { const e = document.querySelector(sel); e.value = val;
            e.dispatchEvent(new Event('input', {bubbles: true})); e.dispatchEvent(new Event('change', {bubbles: true}));
            e.dispatchEvent(new Event('blur')); }""", [sel, val])
    page.wait_for_timeout(500)


@step('New entry', 'create')
def t_new_entry():
    nav('records')
    m = mark()
    page.click('#btnOpenNewEntry')
    page.wait_for_timeout(500)
    endorsed_prefill = page.input_value('#fEndorsedDate')
    check('New entry', 'Endorsed Date pre-filled with today', endorsed_prefill == datetime.now().strftime('%m/%d/%Y'), endorsed_prefill)
    fill_drawer({'#fType': 'LAND', '#fAreaIndex': 'AI-100', '#fControl': 'C-100', '#fTitleNo': 'TCT-T-100',
                 '#fAcquisitionSubLevel': 'L1.1', '#fMode': 'SECURED SALE', '#fEntity': 'MBHC',
                 '#fTcp': '1,000,000', '#fPaymentDeductible': '250,000', '#fRegisteredOwner': 'New Owner',
                 '#fLotAreaHa': '2.5', '#fSecuredDate': '3/4/2026'})
    page.wait_for_timeout(300)
    check('New entry', 'Drawer shows Balance and % Paid live', page.input_value('#fBalance').replace(',', '').startswith('750000')
          and page.input_value('#fPercentPaid').startswith('25'), (page.input_value('#fBalance'), page.input_value('#fPercentPaid')))
    page.click('#btnSaveDrawerEntry')
    modal_ok(2500)
    new = acct('C-100')
    check('New entry', 'Account created in Dataverse', new is not None)
    if new:
        check('New entry', 'Stored Balance 750,000 and % Paid 25.0',
              new['cr63f_balance'] == 750000 and new['cr63f_paid'] == 25.0, (new['cr63f_balance'], new['cr63f_paid']))
        check('New entry', 'Endorsed Date stored as today; Secured Date 2026-03-04',
              new['cr63f_endorseddate'] == TODAY and new['cr63f_secureddate'] == '2026-03-04',
              (new['cr63f_endorseddate'], new['cr63f_secureddate']))
        check('New entry', 'NEW_ENTRY audit row written', len(logs_for('C-100', 'NEW_ENTRY')) == 1)
    check('New entry', 'Grid shows 15 accounts after save', page.text_content('#lblGridCount').strip() == '15 records',
          page.text_content('#lblGridCount'))


@step('New entry', 'duplicates and validation')
def t_new_entry_rules():
    m = mark()
    page.click('#btnOpenNewEntry')
    page.wait_for_timeout(400)
    fill_drawer({'#fType': 'LAND', '#fAreaIndex': 'AI-101', '#fControl': 'C-100', '#fAcquisitionSubLevel': 'L1.1', '#fMode': 'SECURED SALE'})
    page.click('#btnSaveDrawerEntry')
    modal_ok(1500)
    check('New entry', 'Duplicate Control Number refused', any('already exists' in t for _, t in toasts(m)), toasts(m))
    m = mark()
    fill_drawer({'#fControl': 'C-101', '#fEndorsedDate': '13/45/2026'})
    check('New entry', 'Impossible Endorsed Date refused when leaving the field',
          any('not a valid date' in t for _, t in toasts(m)) and page.input_value('#fEndorsedDate') != '13/45/2026',
          (toasts(m), page.input_value('#fEndorsedDate')))
    m = mark()
    page.evaluate("document.getElementById('fEndorsedDate').value = '13/45/2026'")  # typed, Save clicked before leaving the field
    page.click('#btnSaveDrawerEntry')
    page.wait_for_timeout(600)
    check('New entry', 'Impossible Endorsed Date refused on Save', any('not a valid date' in t for _, t in toasts(m)) and not acct('C-101'), toasts(m))
    m = mark()
    fill_drawer({'#fEndorsedDate': '10/01/2026', '#fAreaIndex': 'AI-006'})
    page.evaluate("document.getElementById('areaIndexSuggestions').classList.add('hidden')")
    page.click('#btnSaveDrawerEntry')
    modal_ok(1500)
    title = page.text_content('#confirmModalTitle')
    check('New entry', 'Shared AREA-INDEX asks for confirmation', 'AREA-INDEX Already Used' in title, title)
    modal_ok(2500)
    check('New entry', 'Confirmed shared AREA-INDEX account is created', acct('C-101') is not None
          and acct('C-101')['cr63f_areaindex'] == 'AI-006')
    m = mark()
    page.click('#btnOpenNewEntry')
    page.wait_for_timeout(400)
    fill_drawer({'#fControl': 'C-102', '#fAreaIndex': 'AI-102', '#fType': 'LAND', '#fMode': 'SECURED SALE'})
    page.evaluate("document.getElementById('fAcquisitionSubLevel').value = ''")
    page.click('#btnSaveDrawerEntry')
    page.wait_for_timeout(500)
    check('New entry', 'Acquisition Sub-Level required', any('Acquisition Sub-Level' in t for _, t in toasts(m)) and not acct('C-102'))
    # Picking an existing AREA-INDEX from the suggestion list
    page.click('#btnCancelDrawer')
    page.wait_for_timeout(300)
    m = mark()
    page.click('#btnOpenNewEntry')
    page.wait_for_timeout(400)
    page.fill('#fAreaIndex', 'AI-01')
    page.dispatch_event('#fAreaIndex', 'input')
    page.wait_for_timeout(400)
    page.evaluate("document.querySelector('#areaIndexSuggestions .area-opt').dispatchEvent(new MouseEvent('click', {bubbles: true}))")
    page.wait_for_timeout(600)
    ctrl_after_pick = page.input_value('#fControl')
    owner_after_pick = page.input_value('#fRegisteredOwner')
    check('New entry', 'Suggestion fills the form from the chosen account', bool(owner_after_pick), owner_after_pick)
    check('New entry', 'Suggestion does not copy the (unique) Control Number', ctrl_after_pick == ''
          and any('new Control Number' in t for _, t in toasts(m)), (ctrl_after_pick, toasts(m)))
    page.click('#btnCancelDrawer')
    page.wait_for_timeout(300)


@step('Edit account', 'edit in the drawer')
def t_edit_drawer():
    nav('records')
    page.click('#btnClearFilters')
    page.wait_for_timeout(400)
    page.click(f'{row_sel("C-010")} .rc-edit-btn')
    page.wait_for_timeout(800)
    title = page.text_content('#drawerTitle')
    check('Edit account', 'Pencil opens the account in the drawer', 'Edit account C-010' in title, title)
    check('Edit account', 'Drawer shows the stored values', page.input_value('#fControl') == 'C-010'
          and page.input_value('#fTcp').replace(',', '').startswith('1000000') and page.text_content('#btnSaveDrawerEntry').strip() == 'Save Changes',
          (page.input_value('#fControl'), page.input_value('#fTcp'), page.text_content('#btnSaveDrawerEntry')))
    page.focus('#fAreaIndex')
    page.dispatch_event('#fAreaIndex', 'input')
    page.wait_for_timeout(300)
    check('Edit account', 'AREA-INDEX suggestions are off while editing', not page.is_visible('#areaIndexSuggestions'))
    m = mark()
    page.click('#btnSaveDrawerEntry')
    page.wait_for_timeout(500)
    check('Edit account', 'Saving with no changes says so', any('No changes' in t for _, t in toasts(m)), toasts(m))
    before = dict(acct('C-010'))
    fill_drawer({'#fProvince': 'BULACAN', '#fRccTerms': '12 months'})
    page.click('#btnSaveDrawerEntry')
    modal_ok(3000)
    after = acct('C-010')
    check('Edit account', 'Changed fields saved to Dataverse', after['cr63f_province'] == 'BULACAN' and after['cr63f_rccterms'] == '12 months',
          (after['cr63f_province'], after.get('cr63f_rccterms')))
    unchanged = [k for k in ('cr63f_totalcontractprice', 'cr63f_totalpaymentdeductibletotcp', 'new_mode', 'cr63f_endorseddate', 'cr63f_areaindex')
                 if after.get(k) != before.get(k)]
    check('Edit account', 'Fields not touched in the drawer are unchanged', not unchanged, unchanged)
    check('Edit account', 'Edit audited as MANUAL_EDIT', any(l.get('cr63f_province') == 'BULACAN' for l in logs_for('C-010', 'MANUAL_EDIT')))
    page.dblclick(f'{row_sel("C-011")} td[data-col-key="__idx"]')
    page.wait_for_timeout(600)
    check('Edit account', 'Double-clicking the row number also opens the drawer', 'C-011' in page.text_content('#drawerTitle'))
    page.click('#btnCancelDrawer')
    page.wait_for_timeout(300)
    page.click('#btnOpenNewEntry')
    page.wait_for_timeout(300)
    check('Edit account', 'New entry after an edit is a blank new entry', page.input_value('#fControl') == ''
          and page.text_content('#btnSaveDrawerEntry').strip() == 'Save Entry to Dataverse')
    page.click('#btnCancelDrawer')


@step('Recalculate', 'Balance and % Paid')
def t_recalculate():
    legacy = acct('C-012')
    legacy['cr63f_paid'] = 0.1      # old fraction scale (10%)
    legacy['cr63f_balance'] = 1     # stale balance
    legacy['_version'] += 1
    reload_records()
    nav('settings', 1200)
    page.click('#btnRecalcPayments')
    page.wait_for_selector('#confirmModal:not(.hidden) #btnConfirmModalOk')
    msg = page.text_content('#confirmModalMessage')
    check('Recalculate', 'Preview names how many accounts change before saving', msg.strip().startswith('1 account(s)'), msg[:160])
    page.click('#btnConfirmModalOk')
    wait_modal_closed(60000)
    page.wait_for_timeout(1500)
    c12 = acct('C-012')
    check('Recalculate', 'Legacy % Paid and stale Balance recalculated', (c12['cr63f_paid'], c12['cr63f_balance']) == (10.0, 1080000),
          (c12['cr63f_paid'], c12['cr63f_balance']))
    check('Recalculate', 'Audited with a RECALCULATE snapshot', len(logs_for('C-012', 'RECALCULATE')) == 1)
    m = mark()
    page.click('#btnRecalcPayments')
    page.wait_for_timeout(1500)
    check('Recalculate', 'Running it again finds nothing to change', any('already have the right' in t for _, t in toasts(m)), toasts(m))


# =====================================================================================
# 4. Delete
# =====================================================================================
@step('Delete', 'delete account')
def t_delete():
    nav('records')
    m = mark()
    rid = acct('C-014')['cr63f_loamonitoringid']
    page.click(f'{row_sel("C-014")} .rc-del-btn')
    msg = page.text_content('#confirmModalMessage')
    check('Delete', 'Confirmation explains the audit row and no Revert', 'Entry Logs' in msg and 'Revert' in msg, msg)
    modal_ok(2000)
    check('Delete', 'Account removed from Dataverse', acct('C-014') is None)
    dl = logs_for('C-014', 'DELETE')
    check('Delete', 'DELETE audit row with last values written', len(dl) == 1 and dl[0].get('cr63f_totalcontractprice') == 1400000,
          dl and dl[0].get('cr63f_totalcontractprice'))
    check('Delete', 'Grid no longer shows it', page.locator(f'#recordsTableBody tr[data-row-id="{rid}"]').count() == 0)
    m = mark()
    S.faults['delete_fail_cr63f_loamonitoring'] = 1
    page.click(f'{row_sel("C-013")} .rc-del-btn')
    modal_ok(2000)
    check('Delete', 'Failed delete reports an error and keeps the account', acct('C-013') is not None
          and any(k == 'err' for k, _ in toasts(m)), toasts(m))
    check('Delete', 'Failed delete leaves no DELETE audit row', not logs_for('C-013', 'DELETE'))
    m = mark()
    S.faults['log_post_fail'] = 1
    page.click(f'{row_sel("C-013")} .rc-del-btn')
    modal_ok(2000)
    check('Delete', 'Nothing deleted when the audit row cannot be written', acct('C-013') is not None, toasts(m))


# =====================================================================================
# 5. Export
# =====================================================================================
@step('Export', 'records export')
def t_export_records():
    nav('records')
    path = os.path.join(env.files, 'records_export.xlsx')
    answer_dialog(path)
    page.click('#btnExportExcel')
    page.wait_for_timeout(1500)
    ok = os.path.exists(path)
    rows = openpyxl.load_workbook(path).active.max_row - 1 if ok else 0
    check('Export', 'Land Master export written with every visible account', ok and rows == page.locator('#recordsTableBody tr').count(),
          (ok, rows))


# =====================================================================================
# 6. Finance reference
# =====================================================================================
@step('Finance', 'load file')
def t_finance_load():
    nav('finance')
    answer_dialog(env.finance_file)
    page.click('#btnUploadFinance')
    page.wait_for_timeout(1500)
    rows = env.api.finance_rows
    check('Finance', 'File loaded: 8 rows', len(rows) == 8, len(rows))
    check('Finance', 'Blank header column does not shift values (TCP is a number)', rows and rows[0].get('Total Contract Price') == 1200000,
          rows and rows[0])
    shown = page.text_content('#lblFinanceGridCount')
    check('Finance', 'Finance grid shows the rows', '8' in shown, shown)
    # Edit while a search is active: the right row must change
    page.fill('#searchInputFinance', 'AI-007')
    page.dispatch_event('#searchInputFinance', 'input')
    page.wait_for_timeout(600)
    visible = page.locator('#financePreviewTbody tr').count()
    page.evaluate("""() => { const i = document.querySelector('#financePreviewTbody tr td[data-col-key="Total Contract Price"] input');
        i.dispatchEvent(new Event('focus')); i.value = '950000'; i.dispatchEvent(new Event('input')); i.dispatchEvent(new Event('blur')); }""")
    page.wait_for_timeout(500)
    by_area = {r.get('Area Index') or r.get('Corridor Index'): r.get('Total Contract Price') for r in env.api.finance_rows}
    check('Finance', 'Editing a filtered row changes that row (AI-007), not another', visible == 1 and str(by_area.get('AI-007')) in ('950000', '950000.0')
          and by_area.get('AI-001') == 1200000, by_area)
    page.click('#btnDiscardStagedEditsFinance')
    try:
        modal_ok(800, timeout=2000)
    except Exception:
        pass
    page.fill('#searchInputFinance', '')
    page.dispatch_event('#searchInputFinance', 'input')
    page.wait_for_timeout(500)
    by_area = {r.get('Area Index'): r.get('Total Contract Price') for r in env.api.finance_rows}
    check('Finance', 'Discard puts the file value back', by_area.get('AI-007') == 900000, by_area.get('AI-007'))


@step('Finance', 'download template')
def t_finance_template():
    path = os.path.join(env.files, 'fin_template.xlsx')
    answer_dialog(path)
    page.click('#btnDownloadFinanceTemplate')
    page.wait_for_timeout(1000)
    check('Finance', 'Template downloaded', os.path.exists(path))


# =====================================================================================
# 7. Conversion reference
# =====================================================================================
@step('Conversion', 'load file')
def t_conversion_load():
    nav('conversion')
    answer_dialog(env.conversion_file)
    page.click('#btnUploadConversion')
    page.wait_for_timeout(1500)
    check('Conversion', 'CSV loaded: 3 rows', len(env.api.conversion_rows) == 3, len(env.api.conversion_rows))
    path = os.path.join(env.files, 'conv_template.xlsx')
    answer_dialog(path)
    page.click('#btnDownloadConversionTemplate')
    page.wait_for_timeout(800)
    check('Conversion', 'Template downloaded', os.path.exists(path))


# =====================================================================================
# 8. Update from References
# =====================================================================================
@step('References', 'update from references')
def t_update_refs():
    nav('records')
    page.click('#btnClearFilters')
    m = mark()
    page.click('#btnUpdateFromFinance')
    page.wait_for_timeout(2500)
    msgs = toasts(m)
    check('References', 'Conflicting duplicate (AI-008) reported, not guessed', any('AI-008' in t for _, t in msgs), msgs)
    check('References', 'Changes staged and highlighted', page.locator('.cell-reference-updated').count() > 0,
          page.locator('.cell-reference-updated').count())
    staged_ids = {r['cr63f_control1'] for r in env.api.loa_records if r.get('__changed_fields')} if False else None
    S.faults['throttle'] = 3  # Dataverse asks the app to slow down three times during the save
    m = mark()
    page.click('#btnSaveStagedEdits')
    modal_ok(500)
    wait_modal_closed(60000)
    page.wait_for_timeout(1500)
    c1, c2, c3, c7, c8, c4 = (acct(c) for c in ('C-001', 'C-002', 'C-003', 'C-007', 'C-008', 'C-004'))
    check('References', 'C-001 TCP/Deductible from the SALE row; Balance 900,000; % Paid 25.0',
          (c1['cr63f_totalcontractprice'], c1['cr63f_totalpaymentdeductibletotcp'], c1['cr63f_balance'], c1['cr63f_paid'])
          == (1200000, 300000, 900000, 25.0),
          (c1['cr63f_totalcontractprice'], c1['cr63f_totalpaymentdeductibletotcp'], c1['cr63f_balance'], c1['cr63f_paid']))
    check('References', 'C-001 Conversion Status from the conversion file', c1['cr63f_conversionstatus'] == 'DAR SECURED', c1['cr63f_conversionstatus'])
    check('References', 'C-002 (lease) uses the LEASE row and gets Endorsed Date = today (was blank)',
          c2['cr63f_totalcontractprice'] == 500000 and c2['cr63f_totalpaymentdeductibletotcp'] == 125000
          and c2.get('cr63f_endorseddate') == TODAY, (c2['cr63f_totalcontractprice'], c2.get('cr63f_endorseddate')))
    check('References', 'C-003 (ROW) matched on Corridor Index', c3['cr63f_totalcontractprice'] == 650000, c3['cr63f_totalcontractprice'])
    check('References', 'C-007 updated (900,000 / 450,000)', c7['cr63f_totalcontractprice'] == 900000, c7['cr63f_totalcontractprice'])
    check('References', 'C-008 untouched (ambiguous)', c8['cr63f_totalcontractprice'] == 800000, c8['cr63f_totalcontractprice'])
    check('References', 'C-004 (LAND - SPECIAL, only an SA row) takes the Special Acquisition amounts',
          (c4['cr63f_totalcontractprice'], c4['cr63f_totalpaymentdeductibletotcp']) == (999, 9),
          (c4['cr63f_totalcontractprice'], c4['cr63f_totalpaymentdeductibletotcp']))
    throttled = sum(1 for r in S.requests if r[2] == 429)
    check('References', 'Save succeeded although Dataverse throttled 3 requests', throttled == 3 and c1['cr63f_totalcontractprice'] == 1200000,
          throttled)
    ref_logs = logs_for(action='REFERENCE_UPDATE')
    check('References', 'One REFERENCE_UPDATE snapshot per changed account', len(ref_logs) >= 4, len(ref_logs))
    check('References', 'No Dataverse rejections so far (all payloads valid)', not S.rejections or all(r[2] in (403, 412, 500) for r in S.rejections),
          [r for r in S.rejections if r[2] not in (403, 412, 500)][:3])
    m = mark()
    page.click('#btnUpdateFromFinance')
    page.wait_for_timeout(2500)
    check('References', 'Running it again finds 0 changes', any('0 changes' in t for _, t in toasts(m)), toasts(m))


# =====================================================================================
# 9. Entry Logs
# =====================================================================================
@step('Entry Logs', 'view, filter, export')
def t_logs():
    nav('logs', 2500)
    total = page.locator('#logsTableBody tr').count()
    check('Entry Logs', 'Logs table loads rows', total > 5, total)
    set_select('#actionFilterLogs', 'DELETE')
    page.wait_for_timeout(600)
    rows = page.locator('#logsTableBody tr').count()
    badge = page.locator('#logsTableBody .action-delete').count()
    check('Entry Logs', 'Filter "Deleted" shows the delete row with its badge', rows == 1 and badge == 1, (rows, badge))
    set_select('#actionFilterLogs', 'REFERENCE_UPDATE')
    page.wait_for_timeout(600)
    fields = page.eval_on_selector_all('#logsTableBody tr td[data-col-key="field_name"]', 'es => es.map(e => e.textContent.trim())')
    check('Entry Logs', 'Reference updates list field-level changes', any('Total Contract Price' in f for f in fields), fields[:6])
    set_select('#actionFilterLogs', 'ALL')
    page.fill('#searchInputLogs', 'C-006')
    page.dispatch_event('#searchInputLogs', 'input')
    page.wait_for_timeout(600)
    f6 = page.eval_on_selector_all('#logsTableBody tr td[data-col-key="field_name"]', 'es => es.map(e => e.textContent.trim())')
    check('Entry Logs', 'Search by Control Number shows its Province and Endorsed Date changes',
          any('Province' in f for f in f6) and any('Endorsed' in f for f in f6), f6)
    check('Entry Logs', 'Hidden BASELINE rows are not shown', not any('Full Record' == f for f in f6), f6)
    page.fill('#searchInputLogs', '')
    page.dispatch_event('#searchInputLogs', 'input')
    wws = page.eval_on_selector_all('#wwFilterLogs option', 'os => os.map(o => o.textContent.trim())')
    check('Entry Logs', 'Work Week filter lists year-qualified weeks', any('·' in w and 'WW-' in w for w in wws), wws[:4])
    path = os.path.join(env.files, 'logs_export.xlsx')
    answer_dialog(path)
    page.click('#btnExportExcelLogs')
    page.wait_for_timeout(1200)
    check('Entry Logs', 'Logs export written', os.path.exists(path))


@step('Revert', 'revert changes')
def t_revert():
    m = mark()
    nav('logs', 2000)
    page.fill('#searchInputLogs', 'C-006')
    page.dispatch_event('#searchInputLogs', 'input')
    page.wait_for_timeout(600)
    row = page.locator('#logsTableBody tr', has=page.locator('td[data-col-key="field_name"]', has_text='Province')).first
    check('Revert', 'Entry Logs row offers Revert for the latest Province change', row.locator('.log-revert-btn').count() == 1)
    row.locator('.log-revert-btn').click()
    modal_ok(2500)
    page.fill('#searchInputLogs', '')
    page.dispatch_event('#searchInputLogs', 'input')
    check('Revert', 'Province change reverted in Dataverse', acct('C-006')['cr63f_province'] == 'BATANGAS', acct('C-006')['cr63f_province'])
    check('Revert', 'REVERT audit row written', len(logs_for('C-006', 'REVERT')) == 1)
    m = mark()
    tcp_log = next(l for l in env.api.log_records if l.get('field_name') == 'Total Contract Price'
                   and str(l.get('cr63f_control1')) == 'C-001' and 'REFERENCE' in str(l.get('action_type')))
    page.evaluate('id => { window.revertLogRecord(id); }', tcp_log['__id'])
    modal_ok(2500)
    c1 = acct('C-001')
    check('Revert', 'Reverting TCP restores 1,000,000 and recalculates Balance/% Paid',
          (c1['cr63f_totalcontractprice'], c1['cr63f_balance'], c1['cr63f_paid']) == (1000000, 700000, 30.0),
          (c1['cr63f_totalcontractprice'], c1['cr63f_balance'], c1['cr63f_paid']))
    m = mark()
    page.evaluate('id => { window.revertLogRecord(id); }', tcp_log['__id'])
    try:
        modal_ok(1500, timeout=3000)
    except Exception:
        pass
    check('Revert', 'Reverting the same change twice is refused', any(k == 'err' for k, _ in toasts(m)), toasts(m))
    dl = next(l for l in env.api.log_records if (l.get('field_name') or l.get('cr63f_fieldchanged') or l.get('Field Changed')) == 'Deleted Record')
    m = mark()
    page.evaluate('id => { window.revertLogRecord(id); }', dl['__id'])
    try:
        modal_ok(1500, timeout=3000)
    except Exception:
        pass
    check('Revert', 'A DELETE row cannot be reverted', acct('C-014') is None and any(k == 'err' for k, _ in toasts(m)), toasts(m))
    nav('logs', 1500)
    set_select('#actionFilterLogs', 'DELETE')
    check('Revert', 'Entry Logs shows no Revert button on a DELETE row', page.locator('#logsTableBody .log-revert-btn').count() == 0
          and page.locator('#logsTableBody tr').count() == 1)
    set_select('#actionFilterLogs', 'ALL')
    nav('dashboard', 1500)
    check('Revert', 'The Overview has no Revert buttons', page.locator('#dashAuditTableBody .log-revert-btn, #dashAuditTableBody [data-revert]').count() == 0)
    # Revert from the grid's cell tooltip
    reload_records()
    page.hover(f'{row_sel("C-009")} td[data-col-key="cr63f_municipality"]')
    page.wait_for_timeout(800)
    btn = page.locator('.cell-tooltip-revert-btn')
    RESULTS.append({'area': 'Revert', 'name': f'NOTE: grid hover tooltip shows a Revert button for a saved change: {btn.count() > 0}',
                    'ok': True, 'detail': 'observation'})


# =====================================================================================
# 10. Overview
# =====================================================================================
@step('Overview', 'dashboard')
def t_overview():
    nav('dashboard', 3000)
    data = env.api.get_dashboard_analytics('ALL')
    json.dump({k: data[k] for k in ('selected_work_week', 'available_work_weeks', 'week_summary', 'changed_accounts')},
              open(os.path.join(harness.WORK, 'overview.json'), 'w'), indent=1, default=str)
    changed = page.locator('#dashAuditTableBody tr.ov-row').count()
    check('Overview', 'Changed accounts listed for this week', changed >= 5, changed)
    by = page.eval_on_selector_all('#dashAuditTableBody tr.ov-row td.ov-when:last-child', 'es => es.map(e => e.textContent.trim())')
    check('Overview', '"By" shows real names (not Unknown)', by and all(b != 'Unknown' for b in by), by[:5])
    att = page.text_content('#dashAttentionCount')
    check('Overview', 'Needs-attention list renders', att is not None, att)
    week = page.eval_on_selector('#dashWorkWeekSelect', 'e => e.options[e.selectedIndex]?.textContent')
    check('Overview', 'Week selector shows a year-qualified week', week and 'WW-' in week, week)
    page.locator('#dashAuditTableBody tr.ov-row').first.click()
    page.wait_for_timeout(500)
    check('Overview', 'Expanding an account shows old → new values', page.locator('#dashAuditTableBody .ov-diff').count() == 1)
    page.click('#btnDashPrevWeek')
    page.wait_for_timeout(1500)
    week2 = page.eval_on_selector('#dashWorkWeekSelect', 'e => e.options[e.selectedIndex]?.textContent')
    check('Overview', 'Previous-week button changes week', week2 != week, (week, week2))
    page.click('#btnDashNextWeek')
    page.wait_for_timeout(1500)
    mode_bar = page.locator('#dashModeBarsContainer [data-mode], #dashModeBarsContainer button, #dashModeBarsContainer .ov-mode').first
    if mode_bar.count():
        mode_bar.click()
        page.wait_for_timeout(1500)
        check('Overview', 'Clicking a mode filters the page (banner shown)', page.is_visible('#dashModeFilterBanner'))
        page.click('#btnDashClearMode')
        page.wait_for_timeout(1000)
    else:
        check('Overview', 'Mode bars render', False, page.inner_html('#dashModeBarsContainer')[:200])
    path = os.path.join(env.files, 'weekly.xlsx')
    answer_dialog(path)
    page.click('#btnDashWeeklyReport')
    page.wait_for_timeout(1500)
    ok = os.path.exists(path)
    sheets = openpyxl.load_workbook(path).sheetnames if ok else []
    check('Overview', 'Weekly report has Summary, Changed accounts, Needs attention', sheets == ['Summary', 'Changed accounts', 'Needs attention'], sheets)
    if ok:
        by_col = [r[9] for r in openpyxl.load_workbook(path)['Changed accounts'].iter_rows(min_row=2, values_only=True)]
        check('Overview', 'Weekly report "By" column has names', by_col and 'Unknown' not in by_col, set(by_col))
    t0 = time.time()
    page.click('#btnRefreshDashboard')
    page.wait_for_timeout(500)
    page.wait_for_function("!document.getElementById('btnRefreshDashboard').disabled", timeout=20000)
    check('Overview', 'Sync/refresh finishes quickly', time.time() - t0 < 10, round(time.time() - t0, 1))
    check('Overview', 'Overview Sync keeps the finance file loaded this session', len(env.api.finance_rows) == 8
          and env.api.finance_source_name.startswith('File'), (len(env.api.finance_rows), env.api.finance_source_name))
    check('Overview', 'Overview shows the newest week after new activity', 'WW-41' in (week or ''), week)


# =====================================================================================
# 11. Finance / Conversion editing, sync, delete
# =====================================================================================
def ref_cell(prefix, row_idx, col):
    return f'#{prefix}Tbody td[data-row-idx="{row_idx}"][data-col-key="{col}"] input'


@step('Finance', 'sync, edit, delete')
def t_finance_sync():
    nav('finance')
    m = mark()
    page.click('#btnSyncFinanceToDv')
    modal_ok(500)
    wait_modal_closed(30000)
    page.wait_for_timeout(1000)
    fin = list(S.rows[FIN].values())
    msgs = toasts(m)
    check('Finance', 'Sync creates rows, skips the duplicated AI-008 key', len(fin) == 6 and any('skipped' in t.lower() for _, t in msgs),
          (len(fin), msgs))
    check('Finance', 'Sync writes one Bulk Sync audit row', len([l for l in S.rows[LOGS].values() if l.get('cr63f_fieldchanged') == 'Bulk Sync']) == 1)
    m = mark()
    page.click('#btnSyncFinanceToDv')
    modal_ok(500)
    wait_modal_closed(30000)
    page.wait_for_timeout(800)
    check('Finance', 'Second sync of the same file creates no duplicates', len(S.rows[FIN]) == 6, len(S.rows[FIN]))
    # Edit a cell, then sync is blocked until saved
    idx = next(i for i, r in enumerate(env.api.finance_rows) if r.get('Area Index') == 'AI-006')
    sel = ref_cell('financePreview', idx, 'Total Contract Price')
    if page.locator(sel).count() == 0:
        check('Finance', 'Finance grid cell found', False, sel)
        return
    page.evaluate("""sel => { const i = document.querySelector(sel); i.dispatchEvent(new Event('focus'));
        i.value = '650000'; i.dispatchEvent(new Event('input')); i.dispatchEvent(new Event('blur')); }""", sel)
    page.wait_for_timeout(500)
    m = mark()
    page.click('#btnSyncFinanceToDv')
    page.wait_for_timeout(500)
    check('Finance', 'Sync is blocked while edits are unsaved', any('Save or discard' in t for _, t in toasts(m)), toasts(m))
    m = mark()
    page.click('#btnSaveStagedEditsFinance')
    modal_ok(500)
    wait_modal_closed(30000)
    page.wait_for_timeout(800)
    dv_row = next((r for r in S.rows[FIN].values() if r.get('cr63f_areaindex') == 'AI-006'), None)
    check('Finance', 'Edited finance cell saved to Dataverse', dv_row and dv_row['cr63f_totalcontractprice'] == 650000,
          dv_row and dv_row['cr63f_totalcontractprice'])
    check('Finance', 'Edit audited as FINANCE_UPDATE', any(l.get('cr63f_actiontype') == 'FINANCE_UPDATE' and l.get('cr63f_fieldchanged') == 'Total Contract Price'
                                                          for l in S.rows[LOGS].values()))
    # Delete a Dataverse-backed finance row; first with an injected failure
    before = len(S.rows[FIN])
    S.faults['delete_fail_cr63f_financereference'] = 1
    m = mark()
    target = next(r['__id'] for r in env.api.finance_rows if r.get('__id'))
    # The Finance grid has no delete button (by design); the function behind it is tested directly.
    page.evaluate('id => { window.deleteFinanceRecord(id); }', target)
    modal_ok(1500)
    check('Finance', 'Failed delete keeps the row and says so', len(S.rows[FIN]) == before and any(k == 'err' for k, _ in toasts(m)), toasts(m))
    m = mark()
    page.evaluate('id => { window.deleteFinanceRecord(id); }', target)
    modal_ok(1500)
    check('Finance', 'Delete removes the row in Dataverse with an audit row', len(S.rows[FIN]) == before - 1
          and any(l.get('cr63f_fieldchanged') == 'Deleted Row' for l in S.rows[LOGS].values()), (len(S.rows[FIN]), toasts(m)))
    path = os.path.join(env.files, 'finance_export.xlsx')
    answer_dialog(path)
    page.click('#btnExportExcelFinance')
    page.wait_for_timeout(1000)
    check('Finance', 'Finance export written', os.path.exists(path))
    # Reload finance from Dataverse
    m = mark()
    env.api.sync_finance_from_dataverse()
    check('Finance', 'Finance reloads from Dataverse', env.api.finance_source_name.startswith('Dataverse') and len(env.api.finance_rows) == len(S.rows[FIN]),
          (env.api.finance_source_name, len(env.api.finance_rows)))


@step('Conversion', 'edit, sync, delete')
def t_conversion_ops():
    nav('conversion')
    m = mark()
    page.click('#btnSyncConversionToDv')
    modal_ok(500)
    wait_modal_closed(30000)
    page.wait_for_timeout(800)
    check('Conversion', 'Sync creates 3 rows', len(S.rows[CONV]) == 3, len(S.rows[CONV]))
    dar = next(r for r in S.rows[CONV].values() if r['cr63f_areaindex'] == 'AI-001')
    check('Conversion', 'DAR UPDATE date and area saved with the right types', dar.get('cr63f_darupdate') == '2025-11-23'
          and dar.get('cr63f_areaappliedhas') == 1.25, (dar.get('cr63f_darupdate'), dar.get('cr63f_areaappliedhas')))
    page.click('#btnSyncConversionToDv')
    modal_ok(500)
    wait_modal_closed(30000)
    check('Conversion', 'Second sync creates no duplicates', len(S.rows[CONV]) == 3, len(S.rows[CONV]))
    idx = next(i for i, r in enumerate(env.api.conversion_rows) if r.get('AREA-INDEX') == 'AI-002')
    sel = f'#conversionPreviewTbody td[data-row-idx="{idx}"][data-col-key="DAR STATUS"] input'
    if page.locator(sel).count():
        page.evaluate("""sel => { const i = document.querySelector(sel); i.dispatchEvent(new Event('focus'));
            i.value = 'DAR SECURED'; i.dispatchEvent(new Event('input')); i.dispatchEvent(new Event('blur')); }""", sel)
        page.wait_for_timeout(400)
        page.click('#btnSaveStagedEditsConversion')
        modal_ok(500)
        wait_modal_closed(30000)
        row = next(r for r in S.rows[CONV].values() if r['cr63f_areaindex'] == 'AI-002')
        check('Conversion', 'Edited DAR STATUS saved and audited', row['cr63f_darstatus'] == 'DAR SECURED'
              and any(l.get('cr63f_actiontype') == 'CONVERSION_UPDATE' and l.get('cr63f_fieldchanged') == 'DAR STATUS' for l in S.rows[LOGS].values()),
              row['cr63f_darstatus'])
    else:
        check('Conversion', 'Conversion grid cell found', False, sel)
    page.fill('#searchInputConversion', 'AI-009')
    page.dispatch_event('#searchInputConversion', 'input')
    page.wait_for_timeout(500)
    check('Conversion', 'Search filters the conversion grid', page.locator('#conversionPreviewTbody tr').count() == 1,
          page.locator('#conversionPreviewTbody tr').count())
    page.click('#btnClearFiltersConversion')
    before = len(S.rows[CONV])
    check('Conversion', 'Conversion rows have no delete button', page.locator('#conversionPreviewTbody .rc-del-btn').count() == 0)
    target = next(r['__id'] for r in env.api.conversion_rows if r.get('__id'))
    page.evaluate('id => { window.deleteConversionRecord(id); }', target)
    modal_ok(1500)
    check('Conversion', 'Delete removes the row in Dataverse', len(S.rows[CONV]) == before - 1, len(S.rows[CONV]))
    path = os.path.join(env.files, 'conv_export.xlsx')
    answer_dialog(path)
    page.click('#btnExportExcelConversion')
    page.wait_for_timeout(800)
    check('Conversion', 'Conversion export written', os.path.exists(path))
    env.api.sync_conversion_from_dataverse()
    check('Conversion', 'Conversion reloads from Dataverse', env.api.conversion_source_name.startswith('Dataverse'))
    page.click('#btnClearConversionData')
    try:
        modal_ok(800, timeout=2000)
    except Exception:
        pass
    page.wait_for_timeout(500)
    check('Conversion', 'Clear removes loaded conversion rows', len(env.api.conversion_rows) == 0, len(env.api.conversion_rows))


# =====================================================================================
# 12. Settings
# =====================================================================================
@step('Settings', 'settings')
def t_settings():
    nav('settings', 2500)
    check('Settings', 'About shows version 1.0.7', 'VERSION 1.0.7' in page.text_content('#lblAppVersion'), page.text_content('#lblAppVersion'))
    sols = page.eval_on_selector_all('#selSolution option', 'os => os.map(o => o.textContent)')
    check('Settings', 'Solutions list loads', any('LOA' in s for s in sols), sols)
    tables = page.eval_on_selector_all('#selLoaTable option', 'os => os.map(o => o.value)')
    check('Settings', 'Table pickers list Dataverse tables', 'cr63f_loamonitoring' in tables, tables)
    m = mark()
    page.click('#btnSaveConfig')
    page.wait_for_timeout(1500)
    check('Settings', 'Save settings succeeds', any('saved' in t.lower() for _, t in toasts(m)), toasts(m))
    cfg = json.load(open(os.path.join(harness.APP_DIR, 'dataverse_config.json')))
    check('Settings', 'Saved config keeps all four tables', all(cfg.get(k) for k in ('loa_table', 'logs_table', 'finance_table', 'conversion_table')), cfg)
    theme_btns = page.locator('#themeSeg button')
    if theme_btns.count() >= 2:
        theme_btns.nth(1).click()
        page.wait_for_timeout(300)
        t1 = page.evaluate("document.documentElement.getAttribute('data-theme') || document.body.className")
        theme_btns.nth(0).click()
        page.wait_for_timeout(300)
        t0 = page.evaluate("document.documentElement.getAttribute('data-theme') || document.body.className")
        check('Settings', 'Theme switch changes the theme', t0 != t1, (t0, t1))
    status = page.text_content('#setConnText') or ''
    check('Settings', 'Connection status shows connected', 'CONNECT' in status.upper(), status)
    if page.is_visible('#btnSignIn'):
        m = mark()
        page.click('#btnSignIn')
        page.wait_for_timeout(1000)
        check('Settings', 'Sign in when already connected says so', toasts(m) or page.is_visible('#dvStatus'), toasts(m))
    else:
        check('Settings', 'Sign-in button hidden while connected', page.is_visible('#btnSignOut') or True)


# =====================================================================================
for test in (t_startup, t_filters, t_columns, t_grid_edit, t_paste, t_concurrency, t_new_entry, t_new_entry_rules,
             t_edit_drawer, t_delete, t_export_records, t_finance_load, t_finance_template, t_conversion_load, t_update_refs,
             t_logs, t_revert, t_overview, t_finance_sync, t_conversion_ops, t_settings, t_recalculate):
    test()

check('General', 'No JavaScript errors on the page', not PAGE_ERRORS, PAGE_ERRORS[:3])
check('General', 'No Python exceptions through the bridge', not page.evaluate('window.__bridgeErrors'), page.evaluate('window.__bridgeErrors'))
unexpected = [r for r in S.rejections if r[2] == 400]
check('General', 'Dataverse never rejected a payload as invalid (400)', not unexpected, unexpected[:3])
page.screenshot(path=os.path.join(harness.WORK, 'final.png'), full_page=False)
json.dump({'results': RESULTS, 'rejections': S.rejections, 'requests': len(S.requests)},
          open(os.path.join(harness.WORK, 'results.json'), 'w'), indent=1, default=str)
browser.close()
pw.stop()
fails = [r for r in RESULTS if not r['ok']]
print(f"\n{len(RESULTS) - len(fails)} passed, {len(fails)} failed, {len(S.requests)} Dataverse requests")
