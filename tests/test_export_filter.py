"""Unit tests for Excel export column filtering."""

import os
import unittest
import tempfile
from unittest.mock import patch
from app import Api
import openpyxl


class TestExportColumnFiltering(unittest.TestCase):
    def setUp(self):
        with patch('dataverse.DataverseClient.try_silent_token', return_value=None):
            self.app = Api()
        self.temp_dir = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_export_excludes_system_and_value_columns(self):
        records = [
            {
                '__id': 'guid-1234',
                '__etag': 'etag-abc',
                '_owninguser_value': 'user-guid-001',
                '_transactioncurrencyid_value': 'curr-guid-002',
                '_modifiedonbehalfby_value': 'mod-guid-003',
                '_createdonbehalfby_value': 'cre-guid-004',
                '_modifiedby_value': 'modby-guid-005',
                '_owningbusinessunit_value': 'bu-guid-006',
                '_createdby_value': 'creby-guid-007',
                '_owningteam_value': 'team-guid-008',
                '_ownerid_value': 'owner-guid-009',
                'createdon': '2026-01-01',
                'modifiedon': '2026-01-02',
                'statuscode': 1,
                'statecode': 0,
                'Corridor Index': 'C-100',
                'Total Contract Price': '1,500,000.00',
                'Total Payment Deductible': '500,000.00',
                '% Paid': '33.33%'
            }
        ]
        out_path = os.path.join(self.temp_dir.name, "test_export.xlsx")
        res = self.app.export_records_excel(records, destination_path=out_path)
        self.assertTrue(res.get("success"))

        wb = openpyxl.load_workbook(out_path)
        sheet = wb.active
        headers = [cell.value for cell in sheet[1]]

        # Verify excluded columns are not present
        forbidden = [
            '_owninguser_value', '_transactioncurrencyid_value',
            '_modifiedonbehalfby_value', '_createdonbehalfby_value',
            '_modifiedby_value', '_owningbusinessunit_value',
            '_createdby_value', '_owningteam_value', '_ownerid_value',
            '__id', '__etag', 'createdon', 'modifiedon', 'statuscode', 'statecode'
        ]
        for f in forbidden:
            self.assertNotIn(f, headers)

        # Verify expected custom columns are present
        self.assertIn('Corridor Index', headers)
        self.assertIn('Total Contract Price', headers)
        self.assertIn('Total Payment Deductible', headers)
        self.assertIn('% Paid', headers)

    def test_export_with_explicit_columns_filters_system(self):
        records = [
            {
                'Corridor Index': 'C-100',
                'Total Contract Price': '1,500,000.00',
                '_ownerid_value': 'owner-123',
                '#': 1
            }
        ]
        cols = ['#', 'Corridor Index', '_ownerid_value', 'Total Contract Price']
        out_path = os.path.join(self.temp_dir.name, "test_export_explicit.xlsx")
        res = self.app.export_records_excel(records, columns=cols, destination_path=out_path)
        self.assertTrue(res.get("success"))

        wb = openpyxl.load_workbook(out_path)
        sheet = wb.active
        headers = [cell.value for cell in sheet[1]]

        self.assertNotIn('#', headers)
        self.assertNotIn('_ownerid_value', headers)
        self.assertEqual(headers, ['Corridor Index', 'Total Contract Price'])


if __name__ == "__main__":
    unittest.main()


class TestExportByKey(unittest.TestCase):
    """2026-10-15: the Land Master export reads values by Dataverse name, not by grid label."""

    def setUp(self):
        with patch('dataverse.DataverseClient.try_silent_token', return_value=None):
            self.app = Api()

    def test_columns_whose_label_differs_from_dataverse_name_are_filled(self):
        import openpyxl
        record = {
            '__id': '1', 'cr63f_control1': 'C-1', 'Control': 'C-1',
            'cr63f_totalpaymentdeductibletotcp': '₱500.00', 'Total Payment Deductible to TCP': '₱500.00',
            'cr63f_authorizedrepresentative': 'REP', 'Authorized Representative': 'REP',
            'cr63f_dbadjustedbarangay': 'BRGY', 'DB Adjusted Barangay': 'BRGY',
            'work_week': 'WW-42', 'Work Week': 'WW-42',
        }
        cols = [{'key': 'cr63f_control1', 'label': 'Control Number'},
                {'key': 'cr63f_authorizedrepresentative', 'label': 'Authorized Rep'},
                {'key': 'cr63f_dbadjustedbarangay', 'label': 'Adjusted Barangay'},
                {'key': 'work_week', 'label': 'Work Week'},
                {'key': 'cr63f_totalpaymentdeductibletotcp', 'label': 'Payment Deductible'}]
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, 'x.xlsx')
            res = self.app.export_records_excel([record], columns=cols, destination_path=out)
            self.assertTrue(res['success'], res)
            ws = openpyxl.load_workbook(out).active
            self.assertEqual([c.value for c in ws[1]],
                             ['Control Number', 'Authorized Rep', 'Adjusted Barangay', 'Work Week', 'Payment Deductible'])
            self.assertEqual([c.value for c in ws[2]], ['C-1', 'REP', 'BRGY', 'WW-42', '₱500.00'])

    def test_blank_under_dataverse_name_is_not_replaced_by_an_older_alias(self):
        import openpyxl
        record = {'__id': '1', 'cr63f_registeredowner': '', 'Registered Owner': 'OLD OWNER'}
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, 'x.xlsx')
            self.app.export_records_excel([record], columns=[{'key': 'cr63f_registeredowner', 'label': 'Registered Owner'}],
                                          destination_path=out)
            ws = openpyxl.load_workbook(out).active
            self.assertIn(ws.cell(row=2, column=1).value, (None, ''))
