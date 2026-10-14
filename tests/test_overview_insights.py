"""Tests for the Overview v2 weekly insights: deltas, account grouping, needs-attention flags, report export."""

import os
import tempfile
import unittest
from unittest.mock import patch

from app import Api


def _log(acct, field, old, new, when, action='REFERENCE_UPDATE', author='m.santos', **extra):
    d = {'__id': f'{acct}-{field}-{when}', 'action_type': action, 'cr63f_control1': acct, 'area_index': acct.replace('LOA', 'AI'),
         'field_name': field, 'old_value': old, 'new_value': new, 'log_time': when, 'createdon': when, 'author': author}
    d.update(extra)
    return d


class TestOverviewInsights(unittest.TestCase):
    def setUp(self):
        with patch('dataverse.DataverseClient.try_silent_token', return_value=None):
            self.app = Api()
        self.app.loa_records = [
            {'__id': '1', 'cr63f_control1': 'LOA-1', 'cr63f_areaindex': 'AI-1', 'new_mode': 'DIRECT',
             'cr63f_totalcontractprice': '1000000', 'cr63f_totalpaymentdeductibletotcp': '600000', 'cr63f_conversionstatus': 'DAR SECURED'},
            {'__id': '2', 'cr63f_control1': 'LOA-2', 'cr63f_areaindex': 'AI-2', 'new_mode': 'BROKER',
             'cr63f_totalcontractprice': '500000', 'cr63f_totalpaymentdeductibletotcp': '520000', 'cr63f_conversionstatus': 'DAR FOR APPLICATION'},
            {'__id': '3', 'cr63f_control1': 'LOA-3', 'cr63f_areaindex': 'AI-3', 'new_mode': 'DIRECT',
             'cr63f_totalcontractprice': '800000', 'cr63f_totalpaymentdeductibletotcp': '100000', 'cr63f_conversionstatus': 'NOT YET ENDORSED'},
        ]
        # Tuesday-Monday weeks: WW-39 = 2026-09-22..28, WW-40 = 2026-09-29..10-05
        self.app.log_records = [
            _log('LOA-1', 'Total Payment Deductible to TCP', '300000', '400000', '2026-09-24T09:00:00'),
            _log('LOA-1', 'Total Payment Deductible to TCP', '400000', '500000', '2026-10-01T09:00:00'),
            _log('LOA-1', 'Total Payment Deductible to TCP', '500000', '600000', '2026-10-02T09:00:00'),
            _log('LOA-1', 'Conversion Status', 'DAR FOR APPLICATION', 'DAR SECURED', '2026-10-02T09:00:00'),
            _log('LOA-2', 'Total Payment Deductible to TCP', '700000', '520000', '2026-10-02T10:00:00'),
            _log('LOA-2', 'Conversion Status', 'DAR SECURED', 'DAR FOR APPLICATION', '2026-10-02T10:00:00'),
            _log('LOA-3', 'Initial Entry', '', '', '2026-10-05T10:00:00', action='NEW_ENTRY', author='',
                 cr63f_totalcontractprice='800000', cr63f_totalpaymentdeductibletotcp='100000', cr63f_conversionstatus='NOT YET ENDORSED'),
        ]

    def test_week_summary_deltas(self):
        res = self.app.get_dashboard_analytics(work_week='ALL')
        self.assertEqual(res['selected_work_week'], '2026 · WW-40')
        ws = res['week_summary']
        self.assertEqual(ws['previous_week'], '2026 · WW-39')
        self.assertEqual(ws['week_start'], '2026-09-29')
        self.assertEqual(ws['week_end'], '2026-10-05')
        self.assertEqual(ws['accounts_changed'], 3)
        self.assertEqual(ws['new_entries'], 1)
        # LOA-1 net +200k (400k -> 600k), LOA-2 -180k, LOA-3 new +100k
        self.assertAlmostEqual(ws['deductible_delta'], 120000.0)
        self.assertAlmostEqual(ws['tcp_delta'], 800000.0)
        # Balances are clamped per account: LOA-1 600k -> 400k (-200k); LOA-2 stays
        # overpaid (deductible 700k -> 520k on a 500k TCP), so its balance stays 0;
        # LOA-3 is new with 700k outstanding. Net +500k (not tcp_delta - ded_delta).
        self.assertAlmostEqual(ws['balance_delta'], 500000.0)
        self.assertEqual(ws['secured_delta'], 0)  # LOA-1 gained, LOA-2 lost

    def test_changes_grouped_per_account_first_old_last_new(self):
        res = self.app.get_dashboard_analytics(work_week='ALL')
        loa1 = next(a for a in res['changed_accounts'] if a['account'] == 'LOA-1')
        ded = next(c for c in loa1['changes'] if 'Deductible' in c['field_name'])
        self.assertEqual(ded['old_value'], '400000')
        self.assertEqual(ded['new_value'], '600000')
        self.assertEqual(loa1['sources'], ['reference'])
        self.assertTrue(ded['log_id'])
        loa3 = next(a for a in res['changed_accounts'] if a['account'] == 'LOA-3')
        self.assertTrue(loa3['is_new'])
        self.assertEqual(loa3['last_by'], 'Unknown')

    def test_needs_attention_flags(self):
        res = self.app.get_dashboard_analytics(work_week='ALL')
        kinds = {(n['kind'], n['account']) for n in res['needs_attention']}
        self.assertIn(('payment_down', 'LOA-2'), kinds)
        self.assertIn(('status_back', 'LOA-2'), kinds)
        self.assertIn(('over_paid', 'LOA-2'), kinds)
        self.assertNotIn(('payment_down', 'LOA-1'), kinds)

    def test_mode_filter_applies_to_changes(self):
        res = self.app.get_dashboard_analytics(work_week='ALL', mode='DIRECT')
        accts = {a['account'] for a in res['changed_accounts']}
        self.assertEqual(accts, {'LOA-1', 'LOA-3'})

    def test_previous_week_selected(self):
        res = self.app.get_dashboard_analytics(work_week='2026 · WW-39')
        ws = res['week_summary']
        self.assertEqual(ws['accounts_changed'], 1)
        self.assertAlmostEqual(ws['deductible_delta'], 100000.0)

    def test_weekly_report_workbook(self):
        import openpyxl
        data = self.app.get_dashboard_analytics(work_week='ALL')
        path = os.path.join(tempfile.mkdtemp(), 'r.xlsx')
        Api._write_weekly_report(path, data)
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames, ['Summary', 'Changed accounts', 'Needs attention'])
        self.assertGreaterEqual(wb['Changed accounts'].max_row, 4)


if __name__ == '__main__':
    unittest.main()
