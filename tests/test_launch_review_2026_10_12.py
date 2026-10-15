"""Regression tests for the 2026-10-12 pre-launch review (items A1-A7, B1-B8, C)."""

import csv
import os
import tempfile
import threading
import unittest
from unittest.mock import MagicMock, patch

import openpyxl

import app as app_module
import dataverse as dv
from app import Api, _dialog_path, derive_field_diffs_from_snapshots
from loa_engine import read_tabular_file, write_excel_file


def _api(records=None):
    with patch('dataverse.DataverseClient.try_silent_token', return_value=None):
        api = Api()
    api.dv_client = MagicMock()
    api.dv_client.config = {}
    api.dv_client.fetch_records.return_value = []
    api.dv_client.get_entity_attributes.return_value = {'cr63f_control1': {}}
    api.dv_client.save_loa_record.return_value = {'success': True, 'id': 'rec-1', 'etag': 'W/"2"'}
    api.dv_client.save_log_record.return_value = {'success': True, 'id': 'log-1'}
    api.dv_client.current_user_name.return_value = 'me@example.com'
    api._window = None
    api.loa_records = records or []
    api.pending_session_logs = []
    api.log_records = []
    return api


def _account(**extra):
    record = {
        '__id': 'm1', '__etag': 'W/"1"', 'cr63f_control1': 'C-1', 'Control Number': 'C-1',
        'cr63f_areaindex': 'AI-1', 'cr63f_type': 'LAND', 'new_mode': 'SECURED SALE',
        'cr63f_totalcontractprice': '1000', 'cr63f_totalpaymentdeductibletotcp': '100',
    }
    record.update(extra)
    return record


# ---------------------------------------------------------------- A1
class TestBlankHeaderColumn(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()

    def test_csv_values_stay_under_their_own_header(self):
        path = os.path.join(self.dir, 'f.csv')
        with open(path, 'w', newline='') as f:
            csv.writer(f).writerows([['AREA-INDEX', '', 'TCP'], ['NAS-1', 'note', '1000']])
        headers, rows = read_tabular_file(path)
        self.assertEqual(headers, ['AREA-INDEX', 'TCP'])
        self.assertEqual(rows, [{'AREA-INDEX': 'NAS-1', 'TCP': '1000'}])

    def test_xlsx_values_stay_under_their_own_header(self):
        path = os.path.join(self.dir, 'f.xlsx')
        wb = openpyxl.Workbook()
        wb.active.append(['AREA-INDEX', None, 'TCP'])
        wb.active.append(['NAS-1', 'note', 1000])
        wb.save(path)
        headers, rows = read_tabular_file(path)
        self.assertEqual(rows, [{'AREA-INDEX': 'NAS-1', 'TCP': 1000}])
        os.remove(path)  # the workbook was closed, so the file is not locked

    def test_repeated_header_is_not_overwritten(self):
        path = os.path.join(self.dir, 'd.csv')
        with open(path, 'w', newline='') as f:
            csv.writer(f).writerows([['TCP', 'TCP'], ['1', '2']])
        headers, rows = read_tabular_file(path)
        self.assertEqual(rows, [{'TCP': '1', 'TCP (2)': '2'}])

    def test_export_keeps_formula_like_text_as_text(self):
        path = os.path.join(self.dir, 'x.xlsx')
        write_excel_file(path, ['Remark'], [{'Remark': '=HYPERLINK("x")'}])
        cell = openpyxl.load_workbook(path).active['A2']
        self.assertEqual(cell.data_type, 's')


# ---------------------------------------------------------------- A3
def _client(attr_type='Money', table_attrs=None):
    client = dv.DataverseClient.__new__(dv.DataverseClient)
    client.config = {}
    client._datetime_behavior_cache = {'t': {}}
    client.get_entity_info = lambda table: {'entity_set': 'ts', 'primary_id': 'tid', 'primary_name': 'name'}
    client.get_entity_attributes = lambda table: table_attrs or {
        'cr63f_amount': {'logical': 'cr63f_amount', 'display': 'Amount', 'type': attr_type,
                         'can_create': True, 'can_update': True},
    }
    client.patch_with_version = MagicMock(return_value=(204, ''))
    return client


class TestStrictParsing(unittest.TestCase):
    def sent(self, client):
        return client.patch_with_version.call_args[0][2]

    def test_unreadable_money_is_refused_not_blanked(self):
        for value in ('1.2M', 'abc', '12-34'):
            client = _client()
            with self.assertRaises(ValueError):
                client.save_record('t', {'cr63f_amount': value}, record_id='r')
            client.patch_with_version.assert_not_called()

    def test_money_formats_that_are_numbers(self):
        for value, expected in (('(1,234.00)', -1234.0), ('₱ 1 000', 1000.0), ('25.5%', 25.5), (0, 0.0)):
            client = _client()
            client.save_record('t', {'cr63f_amount': value}, record_id='r')
            self.assertEqual(self.sent(client)['cr63f_amount'], expected)

    def test_placeholders_mean_blank(self):
        for value in ('-', 'N/A', ''):
            client = _client()
            client.save_record('t', {'cr63f_amount': value}, record_id='r')
            self.assertIsNone(self.sent(client)['cr63f_amount'])

    def test_unreadable_date_is_refused(self):
        client = _client('DateTime')
        with self.assertRaises(ValueError):
            client.save_record('t', {'cr63f_amount': '15/20/26'}, record_id='r')

    def test_audit_table_stays_lenient(self):
        client = _client()
        client.save_record('t', {'cr63f_amount': 'Previous Snapshot'}, record_id='r', strict=False)
        self.assertIsNone(self.sent(client)['cr63f_amount'])

    def test_update_returns_new_version(self):
        client = _client()
        client.patch_with_version.return_value = (200, 'W/"9"')
        res = client.save_record('t', {'cr63f_amount': '5'}, record_id='r', expected_etag='W/"8"')
        self.assertEqual(res['etag'], 'W/"9"')
        self.assertEqual(client.patch_with_version.call_args.kwargs['expected_etag'], 'W/"8"')
        self.assertEqual(client.patch_with_version.call_args.kwargs['select_attr'], 'tid')


# ---------------------------------------------------------------- A4
class TestTypeAndModeGuardInGrid(unittest.TestCase):
    def test_grid_cannot_blank_mode(self):
        api = _api([_account()])
        res = api.save_staged_records([{'id': 'm1', 'fields': {'new_mode': '', 'Mode': ''}}])
        self.assertFalse(res['success'])
        self.assertIn('Mode cannot be cleared', res['error'])
        api.dv_client.save_loa_record.assert_not_called()

    def test_grid_cannot_blank_type(self):
        api = _api([_account()])
        res = api.save_staged_records([{'id': 'm1', 'fields': {'cr63f_type': ''}}])
        self.assertFalse(res['success'])
        self.assertIn('TYPE cannot be cleared', res['error'])


# ---------------------------------------------------------------- A5
class TestConditionalSaves(unittest.TestCase):
    def test_grid_save_uses_loaded_version_and_keeps_new_one(self):
        api = _api([_account()])
        api._write_baseline_if_missing = MagicMock(return_value=None)
        res = api.save_staged_records([{'id': 'm1', 'fields': {'cr63f_province': 'BATANGAS'}}])
        self.assertTrue(res['success'], res)
        self.assertEqual(api.dv_client.save_loa_record.call_args.kwargs['expected_etag'], 'W/"1"')
        self.assertEqual(api.loa_records[0]['__etag'], 'W/"2"')

    def test_drawer_save_uses_loaded_version(self):
        api = _api([_account()])
        api._write_baseline_if_missing = MagicMock(return_value=None)
        res = api.save_loa_record_entry({'cr63f_province': 'BATANGAS'}, record_id='m1')
        self.assertTrue(res['success'], res)
        self.assertEqual(api.dv_client.save_loa_record.call_args.kwargs['expected_etag'], 'W/"1"')
        self.assertEqual(api.loa_records[0]['__etag'], 'W/"2"')

    def test_change_by_someone_else_is_reported_not_overwritten(self):
        api = _api([_account()])
        api._write_baseline_if_missing = MagicMock(return_value=None)
        api.dv_client.save_loa_record.side_effect = dv.ConcurrencyError('This record was changed in Dataverse by someone else')
        res = api.save_loa_record_entry({'cr63f_province': 'BATANGAS'}, record_id='m1')
        self.assertFalse(res['success'])
        self.assertIn('someone else', res['error'])
        api.dv_client.save_log_record.assert_not_called()

    def test_record_without_version_saves_unconditionally(self):
        api = _api([_account(__etag='')])
        api._write_baseline_if_missing = MagicMock(return_value=None)
        api.save_staged_records([{'id': 'm1', 'fields': {'cr63f_province': 'X'}}])
        self.assertNotIn('expected_etag', api.dv_client.save_loa_record.call_args.kwargs)


# ---------------------------------------------------------------- A6
class TestAuditedDelete(unittest.TestCase):
    def setUp(self):
        self.api = _api([_account()])
        self.api.dv_client.get_entity_info.return_value = {'entity_set': 'cr63f_loamonitorings'}
        self.order = []
        self.api.dv_client.save_log_record.side_effect = lambda log: self.order.append(('log', log)) or {'id': 'log-del'}
        self.api.dv_client.delete.side_effect = lambda *a, **k: self.order.append(('delete', a, k))
        # Delete reads the current row first (2026-10-15).
        self.api.dv_client.get_record.side_effect = lambda table, rid: {**_account(), '__id': rid}

    def test_audit_row_is_written_before_the_delete(self):
        res = self.api.delete_loa_record_entry('m1')
        self.assertTrue(res['success'], res)
        self.assertEqual([entry[0] for entry in self.order], ['log', 'delete'])
        log = self.order[0][1]
        self.assertEqual(log['action_type'], 'DELETE')
        self.assertEqual(log['cr63f_control1'], 'C-1')
        self.assertEqual(self.api.loa_records, [])
        self.assertTrue(any(l.get('field_name') == 'Deleted Record' for l in self.api.log_records))

    def test_nothing_is_deleted_without_its_audit_row(self):
        self.api.dv_client.save_log_record.side_effect = RuntimeError('logs down')
        res = self.api.delete_loa_record_entry('m1')
        self.assertFalse(res['success'])
        self.api.dv_client.delete.assert_not_called()
        self.assertEqual(len(self.api.loa_records), 1)

    def test_failed_delete_removes_its_audit_row(self):
        self.api.dv_client.delete.side_effect = RuntimeError('403')
        res = self.api.delete_loa_record_entry('m1')
        self.assertFalse(res['success'])
        self.api.dv_client.delete_log_record.assert_called_once_with('log-del')
        self.assertEqual(len(self.api.loa_records), 1)

    def test_delete_row_cannot_be_reverted(self):
        self.api.delete_loa_record_entry('m1')
        row = next(l for l in self.api.log_records if l.get('field_name') == 'Deleted Record')
        res = self.api.revert_log_record(row['__id'])
        self.assertFalse(res['success'])
        self.assertIn('cannot be restored', res['error'])

    def test_delete_is_conditional_on_the_current_version(self):
        self.api.dv_client.get_record.side_effect = lambda table, rid: {**_account(), '__id': rid, '__etag': 'W/"7"'}
        self.api.loa_records[0]['__etag'] = 'W/"7"'
        res = self.api.delete_loa_record_entry('m1')
        self.assertTrue(res['success'], res)
        self.assertEqual(self.order[-1][2], {'expected_etag': 'W/"7"'})

    def test_account_changed_since_load_is_not_deleted(self):
        self.api.dv_client.get_record.side_effect = lambda table, rid: {**_account(), '__id': rid, '__etag': 'W/"8"'}
        self.api.loa_records[0]['__etag'] = 'W/"7"'
        res = self.api.delete_loa_record_entry('m1')
        self.assertFalse(res['success'])
        self.assertIn('changed in Dataverse', res['error'])
        self.api.dv_client.save_log_record.assert_not_called()
        self.api.dv_client.delete.assert_not_called()

    def test_delete_row_shows_in_entry_logs(self):
        log = app_module.build_log_entry('DELETE', 'WW-41', '10/12/2026 09:00 AM', 'AI-1', '', '', 'Deleted Record',
                                         'Active', 'Deleted', full_record=_account())
        diffs = derive_field_diffs_from_snapshots([log])
        self.assertEqual([d['field_name'] for d in diffs], ['Deleted Record'])
        self.assertEqual(app_module._audit_source_action(diffs[0]), 'DELETE')


# ---------------------------------------------------------------- A7
class TestThrottling(unittest.TestCase):
    @patch('dataverse.time.sleep')
    def test_429_is_retried_after_the_requested_wait(self, sleep):
        client = dv.DataverseClient.__new__(dv.DataverseClient)
        client._session = MagicMock()
        busy = MagicMock(status_code=429, headers={'Retry-After': '7'})
        ok = MagicMock(status_code=204)
        client._session.post.side_effect = [busy, ok]
        self.assertIs(client._send('post', 'u'), ok)
        sleep.assert_called_once_with(7.0)

    @patch('dataverse.time.sleep')
    def test_retries_stop_after_the_limit(self, sleep):
        client = dv.DataverseClient.__new__(dv.DataverseClient)
        client._session = MagicMock()
        client._session.get.return_value = MagicMock(status_code=429, headers={})
        self.assertEqual(client._send('get', 'u').status_code, 429)
        self.assertEqual(client._session.get.call_count, dv.DataverseClient._THROTTLE_ATTEMPTS)


class TestBulkReferenceSave(unittest.TestCase):
    def make(self):
        api = _api([_account()])
        api.finance_rows = [{'TYPE': 'SALE', 'Area Index': 'AI-1', 'Total Contract Price': '2000',
                             'Total Payment Deductible to TCP': '100'}]
        api.finance_engine.load_rows(api.finance_rows)
        api._logged_account_keys = MagicMock(return_value={'MASTER:m1'})
        return api

    def test_each_account_saves_master_then_audit_with_its_version(self):
        api = self.make()
        calls = []
        api.dv_client.save_loa_record.side_effect = lambda f, **k: calls.append(('master', k)) or {'id': 'm1', 'etag': 'W/"5"'}
        api.dv_client.save_log_record.side_effect = lambda l: calls.append(('log', {})) or {'id': 'log-1'}
        stale_screen_copy = [dict(_account(), cr63f_totalcontractprice='1', __etag='W/"0"')]
        res = api.populate_loa_records_from_references(stale_screen_copy, persist_to_dataverse=True)
        self.assertTrue(res['success'], res)
        self.assertEqual([c for c, _ in calls], ['master', 'log'])
        # The app's current copy (version W/"1") is used, not the screen's stale one.
        self.assertEqual(calls[0][1]['expected_etag'], 'W/"1"')
        self.assertEqual(api.loa_records[0]['__etag'], 'W/"5"')

    def test_failed_audit_row_rolls_back_with_the_new_version(self):
        api = self.make()
        api.dv_client.save_loa_record.side_effect = [{'id': 'm1', 'etag': 'W/"5"'}, {'id': 'm1', 'etag': 'W/"6"'}]
        api.dv_client.save_log_record.side_effect = RuntimeError('log down')
        res = api.populate_loa_records_from_references([_account()], persist_to_dataverse=True)
        self.assertFalse(res['success'])
        self.assertIn('rolled back', res['error'])
        rollback = api.dv_client.save_loa_record.call_args_list[1]
        self.assertEqual(rollback.kwargs['expected_etag'], 'W/"5"')

    def test_progress_starts_before_baselines(self):
        api = self.make()
        api._window = MagicMock()
        api.populate_loa_records_from_references([_account()], persist_to_dataverse=True)
        first = api._window.evaluate_js.call_args_list[0][0][0]
        self.assertIn('Checking audit history', first)


# ---------------------------------------------------------------- B1
class TestReferenceSync(unittest.TestCase):
    def make(self, existing):
        api = _api()
        api.dv_client.fetch_records.return_value = existing
        api.dv_client.save_finance_record.side_effect = lambda payload, record_id=None: {'id': record_id or 'new-id'}
        return api

    def test_resync_updates_instead_of_duplicating(self):
        existing = [{'__id': 'f1', 'TYPE': 'SALE', 'Area Index': 'AI-1', 'Total Contract Price': '₱1,000.00'}]
        api = self.make(existing)
        api.finance_rows = [{'TYPE': 'SALE', 'Area Index': 'AI-1', 'Total Contract Price': 1500},
                            {'TYPE': 'SALE', 'Area Index': 'AI-2', 'Total Contract Price': 900}]
        res = api.sync_uploaded_finance_to_dataverse()
        self.assertTrue(res['success'], res)
        self.assertEqual(res['counts']['updated'], 1)
        self.assertEqual(res['counts']['created'], 1)
        update_call = api.dv_client.save_finance_record.call_args_list[0]
        self.assertEqual(update_call.kwargs['record_id'], 'f1')
        self.assertEqual(update_call.args[0], {'Total Contract Price': 1500})
        self.assertTrue(any(l.get('field_name') == 'Bulk Sync' for l in api.log_records))

    def test_unchanged_rows_are_not_written(self):
        existing = [{'__id': 'f1', 'TYPE': 'SALE', 'Area Index': 'AI-1', 'Total Contract Price': '₱1,000.00', '% Paid': '10.00'}]
        api = self.make(existing)
        api.finance_rows = [{'TYPE': 'SALE', 'Area Index': 'AI-1', 'Total Contract Price': 1000, '% Paid': '10%'}]
        res = api.sync_uploaded_finance_to_dataverse()
        self.assertEqual(res['counts']['unchanged'], 1)
        api.dv_client.save_finance_record.assert_not_called()
        api.dv_client.save_log_record.assert_not_called()

    def test_key_matching_several_rows_is_skipped(self):
        existing = [{'__id': 'f1', 'TYPE': 'SALE', 'Area Index': 'AI-1'},
                    {'__id': 'f2', 'TYPE': 'SALE', 'Area Index': 'AI-1'}]
        api = self.make(existing)
        api.finance_rows = [{'TYPE': 'SALE', 'Area Index': 'AI-1', 'Total Contract Price': 5}]
        res = api.sync_uploaded_finance_to_dataverse()
        self.assertEqual(res['skipped_count'], 1)
        api.dv_client.save_finance_record.assert_not_called()

    def test_failures_are_reported(self):
        api = self.make([])
        api.dv_client.save_finance_record.side_effect = RuntimeError('400')
        api.finance_rows = [{'TYPE': 'SALE', 'Area Index': 'AI-9'}]
        res = api.sync_uploaded_finance_to_dataverse()
        self.assertFalse(res['success'])
        self.assertEqual(res['error_count'], 1)


# ---------------------------------------------------------------- B2
class TestReferenceDelete(unittest.TestCase):
    def test_failed_dataverse_delete_keeps_the_row(self):
        api = _api()
        api.finance_rows = [{'__id': 'f1', 'Area Index': 'AI-1'}]
        api.dv_client.get_entity_info.return_value = {'entity_set': 'x'}
        api.dv_client.delete.side_effect = RuntimeError('403')
        res = api.delete_finance_record('f1')
        self.assertFalse(res['success'])
        self.assertEqual(len(api.finance_rows), 1)
        api.dv_client.delete_log_record.assert_called_once_with('log-1')

    def test_file_only_row_is_removed_locally(self):
        api = _api()
        api.conversion_rows = [{'AREA-INDEX': 'AI-1'}]
        res = api.delete_conversion_record(0)
        self.assertTrue(res['success'])
        self.assertEqual(api.conversion_rows, [])
        api.dv_client.delete.assert_not_called()


# ---------------------------------------------------------------- B3
class TestAuthor(unittest.TestCase):
    def test_lookup_names_are_read(self):
        client = dv.DataverseClient.__new__(dv.DataverseClient)
        client.api_root = 'https://x/api/data/v9.2/'
        client.get_entity_info = lambda t: {'entity_set': 'logs', 'primary_id': 'logid'}
        client.get_entity_attributes = lambda t: {}
        client._headers = lambda with_formatting=True: {}
        response = MagicMock(ok=True, status_code=200)
        response.json.return_value = {'value': [{
            'logid': 'l1', '_createdby_value': 'guid',
            '_createdby_value@OData.Community.Display.V1.FormattedValue': 'Jonah Inguito',
        }]}
        client._session = MagicMock()
        client._session.get.return_value = response
        rows = client.fetch_records('logs')
        self.assertEqual(rows[0]['createdby'], 'Jonah Inguito')

    def test_rows_saved_this_session_show_the_signed_in_user(self):
        api = _api([_account()])
        log = app_module.build_log_entry('MANUAL_EDIT', 'WW-41', '10/12/2026 09:00 AM', 'AI-1', '', '',
                                         'Total Contract Price', '₱1.00', '₱2.00', full_record=_account())
        api.log_records = [log]
        data = api.get_dashboard_analytics('ALL')
        self.assertEqual(data['changed_accounts'][0]['last_by'], 'me@example.com')


# ---------------------------------------------------------------- B7, C
class TestSmallFixes(unittest.TestCase):
    def test_dialog_path_accepts_string_and_tuple(self):
        self.assertEqual(_dialog_path('C:\\out.xlsx'), 'C:\\out.xlsx')
        self.assertEqual(_dialog_path(('C:\\out.xlsx',)), 'C:\\out.xlsx')
        self.assertEqual(_dialog_path(None), '')
        self.assertEqual(_dialog_path(()), '')

    def test_unaudited_cell_update_is_gone(self):
        self.assertFalse(hasattr(Api, 'update_record_cell'))

    def test_history_check_uses_existing_control_columns(self):
        api = _api()
        api.dv_client.get_entity_attributes.return_value = {'cr63f_control1': {}}
        api._write_baseline_if_missing(_account(), '10/12/2026 09:00 AM')
        query = api.dv_client.fetch_records.call_args.kwargs['filter_query']
        self.assertEqual(query, "cr63f_control1 eq 'C-1'")

    def test_second_sign_in_click_reuses_the_waiting_code(self):
        api = _api()
        api.active_flow = {'user_code': 'ABC', 'verification_uri': 'u', 'message': 'm'}
        api._auth_status = {'in_progress': True}
        api.dv_client.try_silent_token.return_value = False
        res = api.start_sign_in()
        self.assertEqual(res['user_code'], 'ABC')
        api.dv_client.begin_device_flow.assert_not_called()

    def test_version_is_reported(self):
        self.assertEqual(_api().get_initial_state()['app_version'], app_module.APP_VERSION)


if __name__ == '__main__':
    unittest.main()


# ---------------------------------------------------------------- found by the end-to-end run (2026-10-12)
class TestEndToEndFindings(unittest.TestCase):
    def test_reload_keeps_a_finance_file_loaded_this_session(self):
        api = _api()
        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, 'f.csv')
        with open(path, 'w', newline='') as f:
            csv.writer(f).writerows([['TYPE', 'Area Index', 'Total Contract Price'], ['SALE', 'AI-1', '5']])
        api._load_finance_file_path(path)
        res = api.sync_finance_from_dataverse(None, True)
        self.assertTrue(res.get('skipped'))
        self.assertEqual(len(api.finance_rows), 1)
        api.dv_client.fetch_records.assert_not_called()

    def test_explicit_sync_still_loads_from_dataverse(self):
        api = _api()
        api.finance_source_name = 'File: old.xlsx'  # e.g. reopened at start, not loaded this session
        api.dv_client.get_entity_attributes.return_value = {}
        res = api.sync_finance_from_dataverse(None, True)
        self.assertFalse(res.get('skipped'))
        self.assertTrue(api.finance_source_name.startswith('Dataverse'))

    def test_reference_edit_after_sync_is_saved(self):
        """The screen sends a position, the row already has a Dataverse ID: the edit must still be written."""
        api = _api()
        api.finance_rows = [{'__id': 'f1', 'Area Index': 'AI-1', 'Total Contract Price': 100}]
        api.dv_client.save_finance_record.return_value = {'success': True, 'id': 'f1'}
        api.dv_client.save_log_record.return_value = {'success': True, 'id': 'log-9'}
        api.update_finance_cell('0', 'Total Contract Price', 250)
        res = api.save_staged_finance_records([{'id': '0', 'fields': {'Total Contract Price': 250}}])
        self.assertTrue(res['success'], res)
        api.dv_client.save_finance_record.assert_called_once_with({'Total Contract Price': 250}, record_id='f1')
        self.assertEqual(res['logged_count'], 1)


# ---------------------------------------------------------------- decisions 2026-10-12 (round 3)
class TestRecalculatePaymentFields(unittest.TestCase):
    def accounts(self):
        return [
            _account(__id='a1', cr63f_control1='C-1', cr63f_totalcontractprice='1000', cr63f_totalpaymentdeductibletotcp='1000',
                     cr63f_balance='0', cr63f_paid='1'),            # legacy fraction: 1 meant 100%
            _account(__id='a2', cr63f_control1='C-2', cr63f_totalcontractprice='1000', cr63f_totalpaymentdeductibletotcp='250',
                     cr63f_balance='750', cr63f_paid='25.0'),       # already right
            _account(__id='a3', cr63f_control1='C-3', cr63f_totalcontractprice='', cr63f_totalpaymentdeductibletotcp='5',
                     cr63f_balance='9', cr63f_paid='3'),            # no TCP: left as it is
            _account(__id='a4', cr63f_control1='C-4', cr63f_totalcontractprice='800', cr63f_totalpaymentdeductibletotcp='200',
                     cr63f_balance='100', cr63f_paid='25.0'),       # stale Balance
        ]

    def test_preview_counts_without_writing(self):
        api = _api(self.accounts())
        res = api.recalculate_payment_fields(False)
        self.assertEqual((res['to_change'], res['unchanged'], res['skipped_no_tcp']), (2, 1, 1))
        api.dv_client.save_loa_record.assert_not_called()

    def test_save_writes_only_balance_and_paid_with_versions_and_audit(self):
        api = _api(self.accounts())
        api._logged_account_keys = MagicMock(return_value=set())
        res = api.recalculate_payment_fields(True)
        self.assertTrue(res['success'], res)
        self.assertEqual(res['saved_count'], 2)
        calls = {c.kwargs['record_id']: c for c in api.dv_client.save_loa_record.call_args_list}
        self.assertEqual(set(calls), {'a1', 'a4'})
        a1_fields = calls['a1'].args[0]
        self.assertEqual(set(a1_fields), {'cr63f_balance', 'cr63f_paid'})
        self.assertEqual((str(a1_fields['cr63f_balance']), str(a1_fields['cr63f_paid'])), ('0', '100.0'))
        self.assertEqual(calls['a1'].kwargs['expected_etag'], 'W/"1"')
        snapshots = [c.args[0] for c in api.dv_client.save_log_record.call_args_list if c.args[0].get('action_type') == 'RECALCULATE']
        self.assertEqual(len(snapshots), 2)
        c3 = next(r for r in api.loa_records if r['__id'] == 'a3')
        self.assertEqual(c3['cr63f_balance'], '9')  # no TCP: untouched

    def test_recalculation_adds_no_field_rows_to_entry_logs(self):
        """Balance and % Paid are calculated fields, so Entry Logs and the Overview stay clean."""
        before = app_module.build_log_entry('MANUAL_EDIT', '', '10/12/2026 09:00 AM', 'AI-1', '', '', 'Full Record',
                                            'Previous Snapshot', 'Updated Snapshot', full_record=_account(cr63f_paid='1'))
        after = app_module.build_log_entry('RECALCULATE', '', '10/12/2026 10:00 AM', 'AI-1', '', '', 'Full Record',
                                           'Previous Snapshot', 'Updated Snapshot', full_record=_account(cr63f_paid='10.0'))
        for row, created in ((before, '2026-10-12T01:00:00Z'), (after, '2026-10-12T02:00:00Z')):
            row['__createdon_raw'] = created
        diffs = derive_field_diffs_from_snapshots([before, after])
        self.assertEqual(app_module._audit_source_action(after), 'RECALCULATE')
        self.assertFalse([d for d in diffs if d.get('action_type') == 'RECALCULATE'])
