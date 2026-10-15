"""LOA Monitoring Engine.

Pure-Python computation and lookup engine implementing the Excel lookup formula:
  =LET(
      Result,
          IF(
              [@TYPE]="ROW",
              XLOOKUP(
                  [@[INDEX NO]],
                  FILTER(APPENDED_FINANCE[Corridor Index], APPENDED_FINANCE[TYPE]="ROW"),
                  FILTER(APPENDED_FINANCE[% Paid], APPENDED_FINANCE[TYPE]="ROW"),
                  ""
              ),
              XLOOKUP(
                  [@[AREA-INDEX]],
                  APPENDED_FINANCE[Area Index],
                  APPENDED_FINANCE[% Paid],
                  ""
              )
          ),
      IF(
          Result <> "",
          Result,
          IF(
              [@[TITLE MOD]] <> "",
              XLOOKUP(
                  [@[TITLE MOD]],
                  APPENDED_FINANCE[Title Mod],
                  APPENDED_FINANCE[% Paid],
                  ""
              ),
              ""
          )
      )
  )

And derived calculation:
  Balance = Total Contract Price - Total Payment Deductible to TCP
"""

from __future__ import annotations

import csv
from datetime import datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation, localcontext
import io
import math
import os
import re
from typing import Any, Dict, List, Optional, Tuple


def norm_str(val: Any) -> str:
    """Normalize string by stripping and standardizing whitespace."""
    if val is None:
        return ""
    s = str(val).strip()
    if s.endswith(".0") and (s[:-2].isdigit() or (s.startswith("-") and s[1:-2].isdigit())):
        s = s[:-2]
    return s


def norm_key(val: Any) -> str:
    """Normalize lookup keys (case-insensitive, collapsed whitespace, punctuation trimmed)."""
    s = norm_str(val).upper()
    # collapse multiple spaces / dashes
    s = re.sub(r'[\s\-_]+', ' ', s).strip()
    return s


def parse_float(val: Any) -> Optional[float]:
    """Parse numeric float from string/number, stripping currency symbols and commas."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        if math.isnan(val) or math.isinf(val):
            return None
        return float(val)
    s = str(val).strip()
    if not s:
        return None
    # Accounting negatives: (1,234.00) or ₱(1,234.00) mean -1234.00
    negative = bool(re.fullmatch(r'[^\d(]*\(.*\)\s*', s))
    # Scientific notation from Excel/CSV, e.g. 1.2E+05
    sci = re.fullmatch(r'\s*[-+]?\d+(?:\.\d+)?[eE][-+]?\d+\s*', s.replace(',', ''))
    if sci:
        try:
            return float(s.replace(',', ''))
        except ValueError:
            return None
    # Remove currency symbols (₱, $, etc.), commas, spaces, parentheses
    clean = re.sub(r'[^\d.\-+]', '', s)
    if not clean or clean in ('-', '+', '.'):
        return None
    try:
        number = float(clean)
    except ValueError:
        return None
    return -abs(number) if negative else number


def format_currency(val: Any) -> str:
    """Format float as currency with commas and 2 decimals."""
    f = parse_float(val)
    if f is None:
        return ""
    return f"{f:,.2f}"


def format_percent(val: Any) -> str:
    """Format percent value safely (e.g. 75 or 75% -> 75.00%). Preserves scale for 0-100% values."""
    if val is None or val == "":
        return ""
    if isinstance(val, str):
        s = val.strip()
        if not s:
            return ""
        if "%" in s:
            f = parse_float(s)
            if f is not None:
                return f"{f:.2f}%"
            return s
        f = parse_float(s)
    else:
        f = parse_float(val)

    if f is None:
        return str(val)
    # Dataverse stores % Paid on a 0-100 scale (0.5 = 0.50%, 45 = 45.00%).
    # Never rescale: treating values below 1 as fractions turned 0.5% into 50%.
    return f"{f:.2f}%"


def calculate_paid_percent(total_contract_price: Any, payment_deductible: Any) -> Optional[Decimal]:
    """Calculate % Paid from the two financial inputs, never from reference % Paid."""
    def decimal_value(value: Any) -> Optional[Decimal]:
        if isinstance(value, Decimal):
            return value
        text = str(value if value is not None else '').strip()
        if re.fullmatch(r'\s*[-+]?\d+(?:\.\d+)?[eE][-+]?\d+\s*', text.replace(',', '')):
            try:
                return Decimal(text.replace(',', ''))
            except InvalidOperation:
                return None
        clean = re.sub(r'[^\d.\-+]', '', text)
        if clean in ('', '-', '+', '.'):
            return None
        try:
            number = Decimal(clean)
        except InvalidOperation:
            return None
        # Accounting negatives: (1,234.00) means -1234.00
        return -abs(number) if re.fullmatch(r'[^\d(]*\(.*\)\s*', text) else number

    tcp = decimal_value(total_contract_price)
    deductible = decimal_value(payment_deductible)
    if tcp is None or deductible is None:
        return None
    if tcp == 0:
        return Decimal(100) if deductible > 0 else Decimal(0)
    with localcontext() as context:
        context.prec = 28
        return (deductible / tcp) * Decimal(100)


# Dataverse cr63f_paid precision (decimal places).
PAID_PERCENT_STEP = Decimal('0.1')


def paid_percent_for_storage(total_contract_price: Any, payment_deductible: Any) -> Optional[Decimal]:
    """% Paid as written to Dataverse: the calculated value rounded to 1 decimal.

    The Dataverse column cr63f_paid has precision 1, so it keeps one decimal
    place (33.33 is stored as 33.3). Writing 2 decimals made every Update from
    References see 33.3 vs 33.33 as a change and save it again. Every save path
    (manual edit, new entry, revert, Update from References) stores this value.
    """
    value = calculate_paid_percent(total_contract_price, payment_deductible)
    if value is None:
        return None
    return value.quantize(PAID_PERCENT_STEP, rounding=ROUND_HALF_UP)


def parse_log_datetime(value: Any) -> Optional[datetime]:
    """Parse an audit timestamp into a naive local datetime for ordering.

    Accepts ISO strings from Dataverse (UTC 'Z' or offsets, converted to local
    time) and the display formats Dataverse returns as FormattedValue, e.g.
    '10/2/2026 9:00 AM'. Comparing these as text puts '10/...' before '9/...'.
    """
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.astimezone().replace(tzinfo=None) if value.tzinfo else value
    s = str(value).strip()
    if not s:
        return None
    try:
        dt = datetime.fromisoformat(s.replace('Z', '+00:00'))
        return dt.astimezone().replace(tzinfo=None) if dt.tzinfo else dt
    except ValueError:
        pass
    for fmt in (
        '%m/%d/%Y %I:%M:%S %p', '%m/%d/%Y %I:%M %p', '%m/%d/%Y %H:%M:%S',
        '%m/%d/%Y %H:%M', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%m/%d/%Y', '%Y-%m-%d',
    ):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


_SECURED_WORDS = re.compile(r'\b(SECURED|EXEMPT|EXEMPTED|APPROVED|COMPLETED|CONVERTED)\b')
_NEGATING_WORDS = re.compile(r'\b(NOT|NO|YET|PENDING|DENIED|REJECTED|REVOKED|CANCELLED|CANCELED|INSPECTION)\b')
_NOT_STARTED_STATUSES = {'', '-', 'UNKNOWN', 'NONE', 'N/A', 'UNSET', 'NOT YET ENDORSED'}
# Ended without progress (withdrawn, cancelled, pulled out): ranked with not started.
_ENDED_STATUSES = {'CANCELED TITLE', 'CANCELLED TITLE', 'WITHDRAWN', 'PULLED-OUT', 'PULLED OUT'}


def conversion_rank(status: Any) -> int:
    """Rank a conversion/DAR status: 3 secured, 2 in progress, 1 not started.

    Whole-word matching plus a negation check, so 'NOT SECURED', 'UNSECURED',
    'DISAPPROVED' and 'NOT YET CONVERTED' are not counted as secured.
    """
    s = re.sub(r'\s+', ' ', str(status or '').strip().upper())
    if s in _NOT_STARTED_STATUSES or s.startswith('NOT YET ENDORSED') or s in _ENDED_STATUSES:
        return 1
    if _SECURED_WORDS.search(s) and not _NEGATING_WORDS.search(s):
        return 3
    return 2


def format_calculated_percent(total_contract_price: Any, payment_deductible: Any) -> str:
    """Display up to ten decimal places without rounding the underlying calculation."""
    value = calculate_paid_percent(total_contract_price, payment_deductible)
    if value is None:
        return ''
    displayed = format(value, '.10f').rstrip('0').rstrip('.')
    if '.' not in displayed:
        displayed += '.00'
    elif len(displayed.split('.', 1)[1]) == 1:
        displayed += '0'
    return displayed + '%'


def get_work_week_info(d: datetime) -> Tuple[str, datetime, datetime]:
    """Calculate Tuesday-to-Monday corporate work week (WW-01 to WW-52/53) and its start/end dates.

    Decision 2026-10-14: weeks run Tuesday to the following Monday (they were
    Wednesday to Tuesday). A week belongs to the year its Saturday falls in, and
    WW-01 is the first Tuesday-to-Monday week whose Saturday is in that year, so
    week numbers are unchanged by the switch; only Tuesdays move to the next week.
    Every year starts at WW-01 and has 52 or 53 weeks.

    Returns (label, week_start_tuesday, week_end_monday).
    """
    w = d.weekday()  # 0=Mon, 1=Tue, 2=Wed, 3=Thu, 4=Fri, 5=Sat, 6=Sun
    days_since_tue = (w - 1) % 7
    start = d - timedelta(days=days_since_tue)
    end = start + timedelta(days=6)

    saturday = start + timedelta(days=4)  # the week's Saturday decides its year
    y = saturday.year

    jan1 = datetime(y, 1, 1, tzinfo=start.tzinfo)
    first_saturday = jan1 + timedelta(days=(5 - jan1.weekday()) % 7)
    ww1_start = first_saturday - timedelta(days=4)

    week_num = int((start.replace(hour=0, minute=0, second=0, microsecond=0) - ww1_start).days // 7) + 1

    return f"WW-{week_num:02d}", start, end


def get_work_week_label(d: datetime) -> str:
    """Year-qualified work week, e.g. '2026 · WW-40'.

    The year is the year the week's Saturday falls in (the same rule that
    assigns the week number), not the calendar year of ``d``. Using the
    calendar year merged Dec 28-31 2027 (WW-01 of 2028) into Jan 2027's WW-01.
    """
    ww, start, _ = get_work_week_info(d)
    return f"{(start + timedelta(days=4)).year} · {ww}"


def get_work_week(date_value: Any) -> str:
    """Derive Tuesday-to-Monday Work Week (e.g. WW-38) from date string or datetime."""
    if not date_value:
        return ""
    if isinstance(date_value, datetime):
        if date_value.year <= 1900:
            return ""
        return get_work_week_info(date_value)[0]
    s = str(date_value).strip()
    if not s or "1900" in s:
        return ""
    s_date = s.split("T")[0].split(" ")[0].strip()
    s_date = s_date.replace('.', '/')
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y"):
        try:
            d = datetime.strptime(s_date, fmt)
            if d.year <= 1900:
                return ""
            return get_work_week_info(d)[0]
        except ValueError:
            continue
    return ""


def clean_title_variants(val: Any) -> List[str]:
    """Generate normalized lookup tokens and cores from Title Number or Title Mod."""
    if val is None:
        return []
    s = str(val).strip()
    if not s:
        return []
    variants = set()

    # 1. Normalized raw
    s_norm = re.sub(r'[\s\-_]+', ' ', s).strip().upper()
    variants.add(s_norm)

    # 2. Extract inner parenthesis like (T-121628) or (LRC)
    parenthesized = re.findall(r'\(([^)]+)\)', s)
    for p in parenthesized:
        p_clean = re.sub(r'^(?:TCT|OCT|CCT|CLOA|EP|KATIBAYAN|TITLE|NO|NUM|NUMBER|T|O|C)\s*[-.:#\s]*', '', p, flags=re.IGNORECASE).strip()
        digits = re.sub(r'[^\d]', '', p_clean).lstrip('0')
        if digits:
            variants.add(digits)
        alnum = re.sub(r'[^A-Z0-9]', '', p.upper())
        if alnum:
            variants.add(alnum)

    # 3. Strip prefix
    s_clean = re.sub(r'^(?:TCT|OCT|CCT|CLOA|EP|KATIBAYAN|TITLE|NO|NUM|NUMBER|T|O|C)\s*[-.:#\s]*', '', s, flags=re.IGNORECASE).strip()
    s_clean2 = re.sub(r'^(?:TCT|OCT|CCT|CLOA|EP|KATIBAYAN|TITLE|NO|NUM|NUMBER|T|O|C)\s*[-.:#\s]*', '', s_clean, flags=re.IGNORECASE).strip()

    # 4. Pure digits (with leading zeros stripped)
    digits = re.sub(r'[^\d]', '', s_clean2).lstrip('0')
    if digits:
        variants.add(digits)

    # 5. Pure alphanumeric
    alnum = re.sub(r'[^A-Z0-9]', '', s.upper())
    if alnum:
        variants.add(alnum)
    alnum2 = re.sub(r'[^A-Z0-9]', '', s_clean2.upper())
    if alnum2:
        variants.add(alnum2)

    # 6. Specific dash segments e.g. 055-2025002708 -> 2025002708
    for part in re.split(r'[\s\-/,()]+', s):
        part_digits = re.sub(r'[^\d]', '', part).lstrip('0')
        if len(part_digits) >= 3:
            variants.add(part_digits)

    return sorted(v for v in variants if v)


def clean_title_mod(val: Any) -> str:
    """Extract stripped numerical/alphanumeric core from Title Number or Title Mod for robust matching."""
    variants = clean_title_variants(val)
    if not variants:
        return ""
    # Deterministic: return longest pure digits if available, else first sorted token
    digits = [v for v in variants if v.isdigit()]
    if digits:
        return max(digits, key=lambda s: (len(s), s))
    return sorted(variants)[0]


_DAR_STATUS_FIELDS = [
    "cr63f_darstatus", "DAR STATUS", "DAR Status", "darstatus",
    "cr63f_conversionstatus", "Conversion Status", "CONVERSION STATUS", "Status"
]


def _agreeing_candidate(rows: List[Dict[str, Any]], signature) -> Optional[Dict[str, Any]]:
    """Pick a row for a duplicated lookup key only when the duplicates agree.

    One row: that row. Several rows with the same values that matter
    (``signature``): the first, since the choice changes nothing. Several rows
    that disagree: None, because picking whichever row Dataverse returned first
    would be a silent guess.
    """
    unique = list({id(row): row for row in rows}.values())
    if not unique:
        return None
    if len(unique) == 1:
        return unique[0]
    signatures = {signature(row) for row in unique}
    return unique[0] if len(signatures) == 1 else None


class ConversionLookupEngine:
    """Fast indexed lookup engine for Land Conversion Reference data (cr63f_landconversionreference)."""

    def __init__(self, conversion_rows: Optional[List[Dict[str, Any]]] = None):
        self.raw_rows: List[Dict[str, Any]] = []
        self.area_index_map: Dict[str, Dict[str, Any]] = {}
        # Every row per Area Index, so duplicates can be detected instead of
        # silently taking whichever row Dataverse returned first.
        self.area_index_rows: Dict[str, List[Dict[str, Any]]] = {}

        if conversion_rows:
            self.load_rows(conversion_rows)

    def load_rows(self, rows: List[Dict[str, Any]]) -> None:
        self.raw_rows = list(rows)
        self.area_index_map.clear()
        self.area_index_rows.clear()

        for r in rows:
            area_idx = norm_key(self._get_field(r, [
                "cr63f_areaindex", "AREA-INDEX", "AREA INDEX", "Area Index", "AreaIndex", "area_index"
            ]))

            if area_idx:
                self.area_index_rows.setdefault(area_idx, []).append(r)
                if area_idx not in self.area_index_map:
                    self.area_index_map[area_idx] = r

    @staticmethod
    def _get_field(row: Dict[str, Any], candidate_names: List[str]) -> Any:
        for name in candidate_names:
            if name in row and row[name] not in (None, ""):
                return row[name]
        lower_map = {k.lower().replace(" ", "").replace("_", "").replace("-", "").replace("(", "").replace(")", ""): v for k, v in row.items()}
        for name in candidate_names:
            k = name.lower().replace(" ", "").replace("_", "").replace("-", "").replace("(", "").replace(")", "")
            if k in lower_map and lower_map[k] not in (None, ""):
                return lower_map[k]
        return ""

    def lookup(
        self,
        area_index_val: Any
    ) -> Dict[str, Any]:
        """Look up conversion data by Area Index only."""
        norm_area = norm_key(area_index_val)

        matched_row: Optional[Dict[str, Any]] = None
        match_type = ""
        match_key = ""

        ambiguous_count = 0

        # Area Index is the sole conversion lookup key.
        if norm_area and norm_area in self.area_index_rows:
            candidates = self.area_index_rows[norm_area]
            chosen = _agreeing_candidate(candidates, lambda row: norm_key(self._get_field(row, _DAR_STATUS_FIELDS)))
            if chosen is not None:
                matched_row = chosen
                match_type = "Area Index"
                match_key = norm_str(area_index_val)
            else:
                ambiguous_count = len(candidates)

        if not matched_row:
            return {
                "matched": False,
                "ambiguous": bool(ambiguous_count),
                "ambiguous_count": ambiguous_count,
                "match_type": "",
                "match_key": "",
                "conversion_status": "",
                "dar_status": "",
                "nia_status": "",
                "conversion_class": "",
                "dar_batch": "",
                "dar_update": "",
                "matched_row": None,
            }

        dar_status = self._get_field(matched_row, _DAR_STATUS_FIELDS)
        nia_status = self._get_field(matched_row, ["cr63f_niastatus", "NIA STATUS", "NIA Status", "niastatus"])
        conv_class = self._get_field(matched_row, ["cr63f_conversionclass", "CONVERSION CLASS", "Conversion Class", "conversionclass"])
        dar_batch = self._get_field(matched_row, ["cr63f_darbatch", "DAR BATCH", "DAR Batch", "darbatch"])
        dar_update = self._get_field(matched_row, ["cr63f_darupdate", "DAR UPDATE", "DAR Update", "darupdate"])

        return {
            "matched": True,
            "match_type": match_type,
            "match_key": match_key,
            "conversion_status": norm_str(dar_status),
            "dar_status": norm_str(dar_status),
            "nia_status": norm_str(nia_status),
            "conversion_class": norm_str(conv_class),
            "dar_batch": norm_str(dar_batch),
            "dar_update": norm_str(dar_update),
            "matched_row": matched_row,
        }


_FINANCE_TYPE_FIELDS = ["TYPE", "Type", "type", "ROW_TYPE", "cr63f_type"]
# Finance TYPEs that never supply a LAND / LAND - SPECIAL account's amounts:
# ROW belongs to ROW accounts; CWA and SA are contract types kept separate.
_NEVER_FOR_LAND_TYPES = {"ROW", "CWA", "SA"}
# SA = Special Acquisition: a LAND - SPECIAL account with no SALE/LEASE row may
# use its SA row (see FinanceLookupEngine.lookup).
_SPECIAL_ACCOUNT_TYPES = {"LAND SPECIAL"}


def finance_mode_class(mode: Any) -> str:
    """'SALE', 'LEASE' or '' from a Land Master Mode.

    SECURED SALE, SS - STP, SS - FROM ROW, SOLD TO PARTNER ... -> SALE
    SECURED LEASE, SL - STP, SL - FROM LAND STP ...            -> LEASE
    Anything else (RESCINDED, TO BE IDENTIFIED, blank)         -> ''
    """
    s = norm_key(mode)
    if re.search(r'\bLEASE\b', s) or re.match(r'SL\b', s):
        return "LEASE"
    if re.search(r'\bSALE\b', s) or re.match(r'SS\b', s) or s == "SOLD TO PARTNER":
        return "SALE"
    return ""


_TITLE_FIELDS = [
    "Title Mod", "TitleMod", "TITLE MOD", "title_mod", "TITLE", "TITLE NO", "Title No", "titleno", "cr63f_titleno"
]
_TCP_FIELDS = [
    "Total Contract Price", "Total Contract Price (Finance)", "TotalContractPrice",
    "TOTAL CONTRACT PRICE", "TOTAL CONTRACT PRICE (FINANCE)", "TCP", "tcp", "cr63f_totalcontractprice"
]
_DEDUCTIBLE_FIELDS = [
    "Total Payment deductible to TCP", "Total Payment Deductible to TCP", "Total Payment Deductible",
    "Total Payment deductible to TCP (Finance)", "Total Payment Deductible to TCP (Finance)",
    "Total Payment Deductible (Finance)", "Payment Deductible (Finance)",
    "Payment Deductible", "Payment Deductible to TCP", "TotalPaymentDeductibleToTCP",
    "TOTAL PAYMENT DEDUCTIBLE TO TCP", "TOTAL PAYMENT DEDUCTIBLE TO TCP (FINANCE)",
    "cr63f_totalpaymentdeductibletotcp"
]


class FinanceLookupEngine:
    """Fast indexed lookup engine for APPENDED_FINANCE reference data aligned with Excel LET/XLOOKUP formulas."""

    def __init__(self, finance_rows: Optional[List[Dict[str, Any]]] = None):
        self.raw_rows: List[Dict[str, Any]] = []
        # Index 1: TYPE == "ROW", key: Corridor Index (first row, kept for callers)
        self.row_corridor_index: Dict[str, Dict[str, Any]] = {}
        # Index 2: Area Index (first row, kept for callers)
        self.area_index: Dict[str, Dict[str, Any]] = {}
        # Every row per key, so conflicting duplicates are detected, not guessed.
        self.row_corridor_rows: Dict[str, List[Dict[str, Any]]] = {}
        self.area_index_rows: Dict[str, List[Dict[str, Any]]] = {}
        # Index 3: Title Mod
        self.title_mod_index: Dict[str, List[Dict[str, Any]]] = {}
        # Title fallback must retain every candidate: the digit-only variants
        # can collide (for example TCT-043-... and TCT-066-... share a suffix).
        self.title_mod_strict_index: Dict[str, List[Dict[str, Any]]] = {}

        if finance_rows:
            self.load_rows(finance_rows)

    def load_rows(self, rows: List[Dict[str, Any]]) -> None:
        self.raw_rows = list(rows)
        self.row_corridor_index.clear()
        self.area_index.clear()
        self.row_corridor_rows.clear()
        self.area_index_rows.clear()
        self.title_mod_index.clear()
        self.title_mod_strict_index.clear()

        for r in rows:
            type_val = norm_key(self._get_field(r, _FINANCE_TYPE_FIELDS))
            corridor_idx = norm_key(self._get_field(r, [
                "Corridor Index", "CorridorIndex", "CORRIDOR INDEX", "INDEX NO", "Index No", "index_no", "cr63f_indexno"
            ]))
            area_idx = norm_key(self._get_field(r, [
                "Area Index", "AreaIndex", "AREA INDEX", "AREA-INDEX", "area_index", "cr63f_areaindex"
            ]))
            title_mod_raw = self._get_field(r, _TITLE_FIELDS)
            title_mod = norm_key(title_mod_raw)
            title_mod_strict = _strict_title_key(title_mod_raw)

            # Index for TYPE="ROW" (Corridor Index)
            if type_val in ("ROW", "") and corridor_idx:
                self.row_corridor_rows.setdefault(corridor_idx, []).append(r)
                if corridor_idx not in self.row_corridor_index:
                    self.row_corridor_index[corridor_idx] = r

            # Index for Area Index
            if area_idx:
                self.area_index_rows.setdefault(area_idx, []).append(r)
                if area_idx not in self.area_index:
                    self.area_index[area_idx] = r

            # Index for Title Mod (both normalized string and stripped core digits)
            if title_mod:
                self.title_mod_index.setdefault(title_mod, []).append(r)
            for title_variant in clean_title_variants(title_mod_raw):
                self.title_mod_index.setdefault(title_variant, []).append(r)
            if title_mod_strict:
                self.title_mod_strict_index.setdefault(title_mod_strict, []).append(r)

    @staticmethod
    def _get_field(row: Dict[str, Any], candidate_names: List[str]) -> Any:
        for name in candidate_names:
            if name in row and row[name] not in (None, ""):
                return row[name]
        # case-insensitive search
        lower_map = {k.lower().replace(" ", "").replace("_", "").replace("-", "").replace("(", "").replace(")", ""): v for k, v in row.items()}
        for name in candidate_names:
            k = name.lower().replace(" ", "").replace("_", "").replace("-", "").replace("(", "").replace(")", "")
            if k in lower_map and lower_map[k] not in (None, ""):
                return lower_map[k]
        return ""

    def _eligible_for_account(self, row: Dict[str, Any], account_is_row: bool,
                              mode_class: str, explicit_row_only: bool) -> bool:
        """May this Finance row supply a Land Master account's amounts? (See ``lookup``.)"""
        row_type = norm_key(self._get_field(row, _FINANCE_TYPE_FIELDS))
        if account_is_row:
            # ROW rows only. Untyped rows (e.g. an Excel file without TYPE) are
            # accepted only on the Corridor Index key, as before.
            return row_type == "ROW" or (not explicit_row_only and row_type == "")
        if row_type in _NEVER_FOR_LAND_TYPES:
            return False
        if mode_class == "SALE" and row_type == "LEASE":
            return False
        if mode_class == "LEASE" and row_type == "SALE":
            return False
        return True

    def _finance_signature(self, row: Dict[str, Any]) -> Tuple[Optional[float], Optional[float]]:
        tcp = parse_float(self._get_field(row, _TCP_FIELDS))
        ded = parse_float(self._get_field(row, _DEDUCTIBLE_FIELDS))
        return (round(tcp, 2) if tcp is not None else None, round(ded, 2) if ded is not None else None)

    def _match_cascade(self, account_is_row: bool, index_no: Any, area_index_val: Any,
                       title_mod_val: Any, is_eligible) -> Tuple[Optional[Dict[str, Any]], str, str, int]:
        """Primary key steps, then the title fallback, using ``is_eligible(row, explicit_row_only)``.

        Returns (matched_row, match_type, match_key, ambiguous_count).
        """
        norm_idx = norm_key(index_no)
        norm_area = norm_key(area_index_val)

        matched_row: Optional[Dict[str, Any]] = None
        match_type = ""
        match_key = ""
        ambiguous_count = 0

        def eligible(rows: List[Dict[str, Any]], explicit_row_only: bool = False) -> List[Dict[str, Any]]:
            return [row for row in rows if is_eligible(row, explicit_row_only)]

        # Step 1: primary keys. Duplicates (same key, same eligible TYPE) are
        # used only when they agree on TCP and Deductible.
        primary_steps = []
        if account_is_row:
            primary_steps.append(("ROW (Corridor Index)", norm_idx, self.row_corridor_rows, index_no, False))
            primary_steps.append(("ROW (Area Index)", norm_area, self.area_index_rows, area_index_val, True))
        else:
            primary_steps.append(("Area Index", norm_area, self.area_index_rows, area_index_val, False))
        for label, key, index, raw_key, explicit_row_only in primary_steps:
            if matched_row or ambiguous_count or not key or key not in index:
                continue
            candidates = eligible(index[key], explicit_row_only)
            if not candidates:
                continue
            chosen = _agreeing_candidate(candidates, self._finance_signature)
            if chosen is not None:
                matched_row = chosen
                match_type = label
                match_key = norm_str(raw_key)
            else:
                ambiguous_count = len(candidates)

        # Step 2: Fallback to Title Mod if no result. A conflicting primary
        # match stops here: falling back to the title would hide the conflict.
        if not matched_row and not ambiguous_count:
            # Match the title after removing only a leading title-type prefix,
            # keeping meaningful leading number segments (043 vs 066).
            strict_key = _strict_title_key(title_mod_val)
            strict_candidates = self.title_mod_strict_index.get(strict_key, [])
            # ROW and non-ROW never borrow each other's finance values just
            # because a title matches (applies in both directions).
            strict_candidates = eligible(strict_candidates, explicit_row_only=account_is_row)
            if strict_candidates:
                chosen = _agreeing_candidate(strict_candidates, self._finance_signature)
                if chosen is not None:
                    matched_row = chosen
                    match_type = "Title Mod (Normalized Exact)"
                    match_key = strict_key
                else:
                    ambiguous_count = len({id(row): row for row in strict_candidates})
            else:
                fallback_candidates: List[Dict[str, Any]] = []
                for variant in clean_title_variants(title_mod_val):
                    fallback_candidates.extend(self.title_mod_index.get(variant, []))
                query_segments = _title_number_segments(title_mod_val)
                fallback_candidates = [
                    candidate for candidate in eligible(fallback_candidates, explicit_row_only=account_is_row)
                    if _title_segments_compatible(
                        query_segments,
                        _title_number_segments(self._get_field(candidate, _TITLE_FIELDS)),
                    )
                ]
                chosen = _agreeing_candidate(fallback_candidates, self._finance_signature)
                if chosen is not None:
                    matched_row = chosen
                    match_type = "Title Mod (Fallback)"
                    match_key = norm_str(title_mod_val)
                elif fallback_candidates:
                    ambiguous_count = len({id(row): row for row in fallback_candidates})
        return matched_row, match_type, match_key, ambiguous_count

    def lookup(
        self,
        type_val: Any,
        index_no: Any,
        area_index_val: Any,
        title_mod_val: Any,
        mode_val: Any = None,
    ) -> Dict[str, Any]:
        """Find the Finance row for one Land Master account (TYPE-aware, decision 2026-10-07).

        Finance TYPE decides which rows an account may use:
        - ROW account: Finance ROW rows only (Corridor Index, then Area Index, then title).
        - Other accounts: SALE row when Mode is a sale (SECURED SALE, SS - ..., SOLD TO
          PARTNER), LEASE row when Mode is a lease (SECURED LEASE, SL - ...); either
          when Mode says neither. ROW, CWA and SA rows are never used for them,
          except that a LAND - SPECIAL account with no SALE/LEASE match uses its
          SA (Special Acquisition) row (decision 2026-10-12).
        A duplicate is the same key AND same TYPE with different amounts: such an
        account is reported as ambiguous instead of taking the first row.
        """
        norm_t = norm_key(type_val)
        account_is_row = norm_t == "ROW"
        mode_class = finance_mode_class(mode_val)

        def regular(row: Dict[str, Any], explicit_row_only: bool) -> bool:
            return self._eligible_for_account(row, account_is_row, mode_class, explicit_row_only)

        matched_row, match_type, match_key, ambiguous_count = self._match_cascade(
            account_is_row, index_no, area_index_val, title_mod_val, regular)

        # Special Acquisition (decision 2026-10-12): a LAND - SPECIAL account
        # with no SALE/LEASE row may use its SA row. SALE/LEASE still come
        # first, so accounts that already match are unchanged; LAND and ROW
        # accounts never use SA rows; Mode does not filter SA rows.
        if not matched_row and not ambiguous_count and norm_t in _SPECIAL_ACCOUNT_TYPES:
            def special_acquisition(row: Dict[str, Any], explicit_row_only: bool) -> bool:
                return norm_key(self._get_field(row, _FINANCE_TYPE_FIELDS)) == "SA"

            matched_row, match_type, match_key, ambiguous_count = self._match_cascade(
                False, index_no, area_index_val, title_mod_val, special_acquisition)
            if matched_row:
                match_type = f"{match_type} (Special Acquisition)"

        if not matched_row:
            return {
                "matched": False,
                "ambiguous": bool(ambiguous_count),
                "ambiguous_count": ambiguous_count,
                "match_type": "",
                "match_key": "",
                "total_contract_price": None,
                "total_payment_deductible": None,
                "balance": None,
                "percent_paid": None,
                "total_contract_price_formatted": "",
                "total_payment_deductible_formatted": "",
                "balance_formatted": "",
                "percent_paid_formatted": "",
            }

        # Extract values supporting all column naming variations
        tcp_raw = self._get_field(matched_row, _TCP_FIELDS)
        deductible_raw = self._get_field(matched_row, _DEDUCTIBLE_FIELDS)
        balance_raw = self._get_field(matched_row, [
            "Balance", "Balance (Finance)", "BALANCE", "BALANCE (FINANCE)", "balance", "cr63f_balance"
        ])

        tcp_num = parse_float(tcp_raw)
        deductible_num = parse_float(deductible_raw)
        balance_num = parse_float(balance_raw)

        # Compute balance (balance is never negative on overpayment)
        if tcp_num is not None and deductible_num is not None:
            computed_balance = max(0.0, tcp_num - deductible_num)
        elif balance_num is not None:
            computed_balance = max(0.0, balance_num)
        else:
            computed_balance = None

        # The reference table's % Paid is ignored. This is display-only; callers
        # must not persist it as part of an Update from References save.
        calculated_pct = calculate_paid_percent(tcp_raw, deductible_raw)
        pct_val = float(calculated_pct) if calculated_pct is not None else None
        pct_formatted = format_calculated_percent(tcp_raw, deductible_raw)

        return {
            "matched": True,
            "match_type": match_type,
            "match_key": match_key,
            "total_contract_price": tcp_num,
            "total_payment_deductible": deductible_num,
            "balance": computed_balance,
            "percent_paid": pct_val,
            "total_contract_price_formatted": format_currency(tcp_num),
            "total_payment_deductible_formatted": format_currency(deductible_num),
            "balance_formatted": format_currency(computed_balance),
            "percent_paid_formatted": pct_formatted,
            "matched_row": matched_row,
        }


def _title_number_segments(value: Any) -> List[str]:
    """Numeric segments of a title after its type prefix: 'TCT-066-2019006115' -> ['66', '2019006115']."""
    key = _strict_title_key(value)
    return [seg.lstrip('0') or '0' for seg in re.findall(r'\d+', key)]


def _title_segments_compatible(query: List[str], candidate: List[str]) -> bool:
    """True unless the two titles carry different numbers in the same position.

    The shorter list must be the tail of the longer one, so '2019006115'
    matches '066-2019006115', but '999-2019006115' does not match
    '066-2019006115' just because the suffix is shared.
    """
    if not query or not candidate:
        return True
    shorter, longer = (query, candidate) if len(query) <= len(candidate) else (candidate, query)
    return longer[-len(shorter):] == shorter


def _strict_title_key(value: Any) -> str:
    """Normalize a title while preserving its distinguishing numeric segments."""
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    text = re.sub(
        r'^(?:TCT|OCT|CCT|CLOA|EP|KATIBAYAN|TITLE|NO|NUM|NUMBER|T|O|C)\s*[-.:#\s]*',
        '', text,
        flags=re.IGNORECASE,
    ).strip()
    return norm_key(text)


def _header_columns(raw_headers) -> List[Tuple[int, str]]:
    """(column position, header) for every non-blank header.

    The position is kept so a blank header column in the middle of a sheet does
    not shift the values of the columns to its right. A repeated header gets a
    ' (2)', ' (3)' ... suffix instead of overwriting the earlier column.
    """
    columns: List[Tuple[int, str]] = []
    seen: Dict[str, int] = {}
    for position, raw in enumerate(raw_headers or []):
        name = str(raw if raw is not None else '').strip()
        if not name:
            continue
        count = seen.get(name.upper(), 0) + 1
        seen[name.upper()] = count
        columns.append((position, name if count == 1 else f"{name} ({count})"))
    return columns


def _rows_from_columns(columns: List[Tuple[int, str]], raw_rows) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    for raw_row in raw_rows:
        raw_row = list(raw_row or [])
        if not any(value not in (None, '') for value in raw_row):
            continue
        row_dict = {}
        for position, header in columns:
            val = raw_row[position] if position < len(raw_row) else ""
            row_dict[header] = "" if val is None else val
        rows.append(row_dict)
    return rows


def read_tabular_file(file_path: str) -> Tuple[List[str], List[Dict[str, Any]]]:
    """Read CSV or Excel (.xlsx/.xlsm) file into headers and row dicts."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    if ext in ('.xlsx', '.xlsm', '.xltx', '.xltm'):
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        try:
            rows_iter = wb.active.iter_rows(values_only=True)
            try:
                raw_headers = next(rows_iter)
            except StopIteration:
                return [], []
            columns = _header_columns(raw_headers)
            return [h for _, h in columns], _rows_from_columns(columns, rows_iter)
        finally:
            # Release the file so Excel can save it while the app is open.
            wb.close()

    elif ext == '.csv':
        for enc in ('utf-8-sig', 'utf-8', 'cp1252', 'latin-1'):
            try:
                with open(file_path, 'r', encoding=enc, newline='') as f:
                    reader = csv.reader(f)
                    try:
                        raw_headers = next(reader)
                    except StopIteration:
                        return [], []
                    columns = _header_columns(raw_headers)
                    return [h for _, h in columns], _rows_from_columns(columns, reader)
            except UnicodeDecodeError:
                continue
        raise ValueError(f"Could not decode CSV file {file_path}")
    else:
        raise ValueError(f"Unsupported file format: {ext}")


def set_text_safe(cell, value: Any) -> None:
    """Write a value without letting text that starts with '=' become a formula."""
    cell.value = value
    if isinstance(value, str) and value.startswith('='):
        cell.data_type = 's'


def write_excel_file(file_path: str, headers: List[str], rows: List[Dict[str, Any]], sheet_name: str = "LOA Records") -> None:
    """Write rows to Excel file with formatted headers and widths."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Header styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A56DB", end_color="1A56DB", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB'),
    )

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, row_data in enumerate(rows, start=2):
        for col_index, h in enumerate(headers, start=1):
            set_text_safe(ws.cell(row=row_index, column=col_index), row_data.get(h, ""))

    # Auto-adjust column widths
    for col_num, h in enumerate(headers, 1):
        max_len = len(str(h))
        for row in range(2, len(rows) + 2):
            cell_val = ws.cell(row=row, column=col_num).value
            if cell_val is not None:
                max_len = max(max_len, len(str(cell_val)))
        col_letter = get_column_letter(col_num)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    wb.save(file_path)

